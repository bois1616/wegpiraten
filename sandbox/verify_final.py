"""Belegt die Datenübernahme des finalen Builds ohne Ausgabe von Personendaten."""

import argparse
import collections
import importlib.util
import json
from datetime import date
from pathlib import Path
from types import ModuleType
from typing import Any

import openpyxl
from loguru import logger

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("--workbook", type=Path, required=True)
parser.add_argument("--numbers-before", type=Path, required=True)
parser.add_argument("--summary", type=Path, required=True)
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]


def module(name: str) -> ModuleType:
    spec = importlib.util.spec_from_file_location(name, root / "sandbox" / f"{name}.py")
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


migrate = module("migrate")
build = module("build")
prepare = module("prepare")
original = migrate.read_clients()
groups, conflicts = migrate.build_person_groups(original)
final = args.workbook
w = openpyxl.load_workbook(final, data_only=True)


def table(sheet: str, key: str) -> list[dict[str, Any]]:
    ws = w[sheet]
    columns = {c.value: c.column for c in ws[3] if c.value}
    records = []
    for cells in ws.iter_rows(min_row=4, values_only=True):
        row = {name: cells[index - 1] for name, index in columns.items()}
        if row.get(key):
            records.append(row)
    return records


persons = table("Kinder", "person_id")
mandates = table("Aufträge", "mandate_id")
cares = table("Betreuungen", "mandate_id")
relations = table("Zuordnung MA", "mandate_id")
reports = table("Berichte", "report_id")
contacts = table("Ansprechpersonen", "contact_person_id")
pby = {r["person_id"]: r for r in persons}
mby = {r["mandate_id"]: r for r in mandates}
cby = {r["contact_person_id"]: r for r in contacts}
bby = {(r["mandate_id"], r["person_id"]): r for r in cares}
numbers = json.loads((root / "sandbox/mandate_numbers.json").read_text())
before = json.loads(args.numbers_before.read_text())
assert all(numbers[k] == v for k, v in before.items())


def normalized(field: str, value: Any) -> Any:
    if field in ("date_of_birth", "start_date", "end_date", "due_date", "completed_date"):
        return date.fromisoformat(value[:10]).isoformat() if isinstance(value, str) and value else migrate.iso(value)
    return migrate.norm(value)


errors = []


def equal(cid: str, field: str, expected: Any, actual: Any) -> None:
    if normalized(field, expected) != normalized(field, actual):
        errors.append(f"{cid}:{field}")


merges = {}
for cid, members in groups.items():
    pid = "C" + cid.lstrip("C")
    p = pby[pid]
    if len(members) > 1:
        merges[pid] = [r["client_id"] for r in members]
    for source in members:
        for field in migrate.PERSON_FIELDS:
            equal(source["client_id"], field, source.get(field), p.get(field))
for source in original:
    cid = source["client_id"]
    mid = numbers[cid]
    old = mby[mid]
    spec = prepare.FOLLOW_UPS.get(cid)
    current = mby[numbers[cid + "+"]] if spec else old
    for field in (
        "service_type_id",
        "service_requester_id",
        "payer_id",
        "tenant_id",
        "application_number",
        "start_date",
        "end_date",
        "allowed_travel_time",
        "allowed_direct_effort",
        "allowed_indirect_effort",
        "allocation",
    ):
        expected = source.get(field)
        if spec and field == "start_date":
            continue
        equal(cid, field, expected, current.get(field))
    if spec:
        equal(cid, "end_date", spec["old_end"], old["end_date"])
        equal(
            cid,
            "application_number",
            spec.get("old_application_number", source.get("application_number")),
            old["application_number"],
        )
        assert old["notes"].startswith("ZU PRÜFEN:") and current["notes"].startswith("ZU PRÜFEN:")
    contact = cby[current["contact_person_id"]]
    for source_field, target_field in (
        ("sr_ap_first_name", "first_name"),
        ("sr_ap_last_name", "last_name"),
        ("service_requester_id", "service_requester_id"),
    ):
        equal(cid, source_field, source.get(source_field), contact.get(target_field))
    if normalized("gender", source.get("sr_ap_gender")) != normalized("gender", contact.get("gender")):
        assert "uneinheitlich" in (contact["notes"] or "")
    if source.get("notes"):
        assert source["notes"] in (current.get("notes") or "")
    candidates = [r for r in cares if r["mandate_id"] == current["mandate_id"]]
    assert len(candidates) == 1
    care = candidates[0]
    equal(cid, "start_date", source.get("start_date"), care["start_date"])
    expected_end = source.get("end_date") if any(source.get(f) is not None for f in migrate.LEAVE_MARKERS) else None
    equal(cid, "end_date", expected_end, care.get("end_date"))
    for field in migrate.CARE_FIELDS:
        equal(cid, field, source.get(field), care.get(field))
assert not errors, errors
original_pairs = {(r["mandate_id"], r["employee_id"]) for r in migrate.read_relations()}
expected_pairs = {(numbers[cid], emp) for cid, emp in original_pairs} | {
    (numbers[cid + "+"], emp) for cid, emp in original_pairs if cid in prepare.FOLLOW_UPS
}
assert expected_pairs == {(r["mandate_id"], r["employee_id"]) for r in relations}
counts = collections.Counter(r["mandate_id"] for r in relations)
assert all(r["role"] == ("P" if counts[r["mandate_id"]] == 1 else None) for r in relations)
assert all(not p.get("family_id") for p in persons)
assert len(reports) == 153 and all(r["notes"].startswith("ZU PRÜFEN:") and r["mandate_id"] in mby for r in reports)
# Handarbeit Bericht für Bericht vergleichen, nicht nur die Anzahl.
hand = openpyxl.load_workbook(root / "sandbox/wegpiraten_datenbank_neu.xlsx", data_only=True)
ws = hand["Berichte"]
headers = [c.value for c in ws[3]]
hand_reports = {
    r[headers.index("report_id")]: dict(zip(headers, r))
    for r in ws.iter_rows(min_row=4, values_only=True)
    if r[headers.index("report_id")]
}
for r in reports:
    h = hand_reports[r["report_id"]]
    for field in ("mandate_id", "due_date", "report_form", "status", "completed_date"):
        equal(r["report_id"], field, h.get(field), r.get(field))
    if h.get("notes"):
        assert str(h["notes"]) in r["notes"]
assert not errors, errors
formula_errors = [f"{ws.title}!{c.coordinate}" for ws in w for row in ws for c in row if c.data_type == "e"]
assert not formula_errors, formula_errors[:30]
fw = openpyxl.load_workbook(final, data_only=False)
build.verify_excel_strict(fw)
checks = [
    dict(
        art=w["Prüfungen"].cell(r, 8).value,
        titel=w["Prüfungen"].cell(r, 10).value,
        anzahl=w["Prüfungen"].cell(r, 12).value,
    )
    for r in range(7, 7 + len(build.CHECKS))
]
assert next(x["anzahl"] for x in checks if x["titel"] == "Übernommene Angabe zu prüfen (Aufträge)") == 6
assert next(x["anzahl"] for x in checks if x["titel"] == "Übernommene Angabe zu prüfen (Berichte)") == 153
summary = {
    "counts": {
        k: len(v)
        for k, v in [
            ("children", persons),
            ("mandates", mandates),
            ("cares", cares),
            ("relations", relations),
            ("reports", reports),
            ("contacts", contacts),
        ]
    },
    "merges": merges,
    "unchanged_numbers": len(before),
    "new_numbers": {k: v for k, v in numbers.items() if k not in before},
    "checks": checks,
    "formula_errors": 0,
    "field_differences": 0,
}
args.summary.write_text(json.dumps(summary, ensure_ascii=False, indent=2))
logger.info(json.dumps({k: v for k, v in summary.items() if k != "checks"}, ensure_ascii=False))

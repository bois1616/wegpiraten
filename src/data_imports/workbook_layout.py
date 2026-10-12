"""Ordnet Eingaben vor Anzeigen und erhält native Excel-Metadaten."""

import re
from copy import deepcopy
from xml.etree import ElementTree as ET

from openpyxl.formula.tokenizer import Tokenizer
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries

from data_imports.workbook_package import (
    NS,
    WorkbookPackage,
    cell_map,
    expand_shared_formulas,
    serialize,
    set_cell,
    tag,
)

_TABLES = {
    "Kinder": "person",
    "Ansprechpersonen": "masterdata_contact_person",
    "Zuordnung MA": "relation_mandate_emp",
    "Berichte": "report",
}
_DESCRIPTIONS = {
    "Bezugspersonen": "Eine Zeile je Bezugsperson: Name, Adresse und Kontakt einmal erfassen. Dieselbe Person kann mehreren Kindern zugeordnet werden. Beziehung und Sorgeberechtigung stehen in Kind-Bezugsperson, nicht hier.",
    "Kind-Bezugsperson": "Eine Zeile verbindet ein Kind mit einer Bezugsperson. Hier die Rolle zum Kind und die Sorgeberechtigung erfassen; mehrere Bezugspersonen je Kind sind möglich. Namen werden angezeigt, Kontaktdaten in Bezugspersonen gepflegt.",
    "Platzierungen": "Eine Zeile je bekanntem Aufenthalt eines Kindes in einer Einrichtung oder Pflegefamilie, mit Adresse, Kontakt und bekannten Datumsgrenzen. Frühere Aufenthalte behalten. Fehlende Angaben offenlassen; die Platzierung bestimmt weder Geschwistergruppe noch Sorgeberechtigung.",
    "Prüfkatalog": "Eine feste Zeile je Prüfung. Kundenkommentare und Bearbeitungsnotizen hier pflegen, damit sie beim Beheben anderer Befunde erhalten bleiben. Prüfungen zeigt die offenen Prüfungen; Fehlerliste zeigt die betroffenen Datensätze mit Zeilenlinks.",
    "Zuordnung MA": "Eine Zeile ordnet einem Auftrag einen Mitarbeiter mit Rolle P (primär) oder S zu. Gültig ab und Gültig bis begrenzen den Zeitraum; leere Grenzen sind offen. Bei einem Wechsel die alte Zeile beenden und eine neue anlegen. Historische Zeilen behalten. Diese Liste steuert Erfassungsbögen und Berichtszuständigkeit.",
    "Berichte": "Eine Zeile je fälligem oder erledigtem Bericht zu einem Auftrag. Auftrag, Berichtsform, Fälligkeit und Status erfassen. Zuständig ist die Person mit Rolle P aus Zuordnung MA am Fälligkeitstag. Berichtsrhythmus und Intervall am Auftrag erzeugen keine Termine; Termine hier einzeln pflegen.",
}
_REFERENCE = re.compile(r"(?P<a>\$?[A-Z]{1,3})(?P<row>\$?\d+)?(?::(?P<b>\$?[A-Z]{1,3})(?P<last>\$?\d+)?)?$")


def remap_reference(reference: str, mapping: dict[int, int]) -> str:
    """Verschiebt Zell-/Spaltenbereiche; nicht zusammenhängende Bereiche brechen ab."""
    match = _REFERENCE.fullmatch(reference)
    if match is None:
        return reference
    if match["row"] is None and match["b"] is None:
        return reference
    first = column_index_from_string(match["a"].lstrip("$"))
    last = column_index_from_string((match["b"] or match["a"]).lstrip("$"))
    moved = [mapping.get(column, column) for column in range(first, last + 1)]
    if max(moved) - min(moved) + 1 != len(moved):
        raise ValueError(f"Spaltenbereich würde auseinanderfallen: {reference}")
    start = ("$" if match["a"].startswith("$") else "") + get_column_letter(min(moved)) + (match["row"] or "")
    if match["b"] is None:
        return start
    end = ("$" if match["b"].startswith("$") else "") + get_column_letter(max(moved)) + (match["last"] or "")
    return start + ":" + end


def remap_formula(expression: str, title: str, mappings: dict[str, dict[int, int]]) -> str:
    """Ändert nur echte A1-Bezüge; Text und strukturierte Tabellennamen bleiben gleich."""
    leading = expression.startswith("=")
    tokens = Tokenizer(expression if leading else "=" + expression).items
    for token in tokens:
        if token.type == "FUNC" and token.subtype == "OPEN" and ":" in token.value:
            reference, _, function = token.value.partition(":")
            token.value = remap_formula(reference, title, mappings) + ":" + function
            continue
        if token.subtype != "RANGE":
            continue
        colon = ":" if token.value.startswith(":") else ""
        value = token.value[len(colon) :]
        prefix, separator, reference = value.rpartition("!")
        target = prefix.strip("'").replace("''", "'") if separator else title
        if target not in mappings:
            continue
        moved = remap_reference(reference if separator else value, mappings[target])
        token.value = colon + (prefix + separator + moved if separator else moved)
    return ("=" if leading else "") + "".join(token.value for token in tokens)


def group_input_columns(package: WorkbookPackage) -> dict[str, dict[int, int]]:
    """Verschiebt bestehende Spalten und alle Formeln, Listen und Formatierungsbereiche."""
    mappings: dict[str, dict[int, int]] = {}
    roots = {title: package.read_sheet(title) for title in package.sheets}
    tables: dict[str, tuple[str, ET.Element]] = {}
    for title, name in _TABLES.items():
        path, table = package.table(name)
        tables[title] = path, table
        left, top, right, _ = range_boundaries(table.get("ref", ""))
        assert left and top and right
        cells = cell_map(roots[title])
        inputs, displays = [], []
        for column in range(left, right + 1):
            cell = cells.get(f"{get_column_letter(column)}{top + 1}")
            (displays if cell is not None and cell.find(tag("f")) is not None else inputs).append(column)
        mapping = {old: new for new, old in enumerate(inputs + displays, left)}
        if any(old != new for old, new in mapping.items()):
            mappings[title] = mapping
    for title, root in roots.items():
        expand_shared_formulas(root)
        for formula in root.iter():
            if formula.tag.rsplit("}", 1)[-1] in ("f", "formula", "formula1", "formula2") and formula.text:
                formula.text = remap_formula(formula.text, title, mappings)
        if title in mappings:
            mapping = mappings[title]
            for row in root.findall("m:sheetData/m:row", NS):
                for cell in row:
                    cell.set("r", remap_reference(cell.get("r", ""), mapping))
                row[:] = sorted(row, key=lambda cell: column_index_from_string(re.sub(r"\d", "", cell.get("r", "A1"))))
            for node in root.iter():
                if node.tag.rsplit("}", 1)[-1] == "sqref" and node.text:
                    node.text = " ".join(remap_reference(ref, mapping) for ref in node.text.split())
                for attribute in ("ref", "sqref", "activeCell", "topLeftCell"):
                    if attribute in node.attrib and node.tag != tag("c"):
                        if node.tag == tag("mergeCell") and node.get("ref", "").endswith("1"):
                            left, top, right, bottom = range_boundaries(node.get("ref", ""))
                            assert left and top and right and bottom
                            moved = [mapping.get(column, column) for column in range(left, right + 1)]
                            node.set(
                                "ref", f"{get_column_letter(min(moved))}{top}:{get_column_letter(max(moved))}{bottom}"
                            )
                        else:
                            node.set(
                                attribute,
                                " ".join(remap_reference(ref, mapping) for ref in node.get(attribute, "").split()),
                            )
            cols = root.find(tag("cols"))
            moved_columns = []
            for col in cols if cols is not None else []:
                for old in range(int(col.get("min", "1")), int(col.get("max", "1")) + 1):
                    new = deepcopy(col)
                    new.set("min", str(mapping.get(old, old)))
                    new.set("max", new.get("min", ""))
                    moved_columns.append(new)
            if cols is not None:
                cols[:] = sorted(moved_columns, key=lambda col: int(col.get("min", "1")))
        package.save_sheet(title, root)
    for title, (path, table) in tables.items():
        columns = table.find(tag("tableColumns"))
        assert columns is not None
        for formula in table.findall(".//m:calculatedColumnFormula", NS):
            if formula.text:
                formula.text = remap_formula(formula.text, title, mappings)
        if title in mappings:
            mapping = mappings[title]
            columns[:] = [
                column for _, column in sorted((mapping[index], column) for index, column in enumerate(columns, 1))
            ]
            for node in table.iter():
                if "ref" in node.attrib:
                    node.set("ref", remap_reference(node.get("ref", ""), mapping))
        package.put_table(path, table)
    # Auch berechnete Spalten anderer Tabellen können auf die verschobenen Blätter zeigen.
    for path, data in list(package.parts.items()):
        if not re.fullmatch(r"xl/tables/table\d+\.xml", path) or path in {item[0] for item in tables.values()}:
            continue
        table = ET.fromstring(data)
        for formula in table.findall(".//m:calculatedColumnFormula", NS):
            if formula.text:
                formula.text = remap_formula(formula.text, "", mappings)
        package.put_table(path, table)
    for name in package.workbook.findall("m:definedNames/m:definedName", NS):
        title = list(package.sheets)[int(name.get("localSheetId", "0"))] if name.get("localSheetId") else ""
        if name.text:
            name.text = remap_formula(name.text, title, mappings)
    return mappings


def describe_sheets(package: WorkbookPackage) -> None:
    """Erklärt Zweck und Pflege der neuen Blätter sowie der datierten Zuordnung."""
    for title, description in _DESCRIPTIONS.items():
        root = package.read_sheet(title)
        set_cell(root, "A1", description, "100")
        row = root.find('m:sheetData/m:row[@r="1"]', NS)
        assert row is not None
        row.set("ht", "62")
        row.set("customHeight", "1")
        if title in ("Bezugspersonen", "Kind-Bezugsperson", "Platzierungen", "Prüfkatalog"):
            for cell in list(row):
                if cell.get("r") == "B1":
                    row.remove(cell)
            merges = root.find(tag("mergeCells"))
            if merges is None:
                merges = ET.Element(tag("mergeCells"))
                data = root.find(tag("sheetData"))
                assert data is not None
                root.insert(list(root).index(data) + 1, merges)
            reference = "A1:H1" if title == "Prüfkatalog" else "A1:F1"
            if not any(node.get("ref") == reference for node in merges):
                ET.SubElement(merges, tag("mergeCell"), {"ref": reference})
            merges.set("count", str(len(merges)))
        package.save_sheet(title, root)
    root = package.read_sheet("Anleitung")
    for cell in root.findall(".//m:c", NS):
        if (
            package.value(cell)
            == "Nur A:C sowie P:Q (Gültig ab / Gültig bis) bearbeiten. D:O enthalten Formeln: hier keine Daten hineinkopieren."
        ):
            set_cell(
                root,
                cell.get("r", ""),
                "Nur A:E bearbeiten: Auftrag, Mitarbeiter, Rolle, Gültig ab und Gültig bis. F:Q enthalten Formeln; hier keine Daten hineinkopieren.",
                cell.get("s", "26"),
            )
    package.save_sheet("Anleitung", root)
    root = package.read_sheet("Fehlerliste")
    for view in root.findall("m:sheetViews/m:sheetView", NS):
        view.attrib.pop("topLeftCell", None)
        for pane in view.findall(tag("pane")):
            pane.set("topLeftCell", "A5")
        for selection in view.findall(tag("selection")):
            selection.set("activeCell", "A5")
            selection.set("sqref", "A5")
    package.save_sheet("Fehlerliste", root)


def finish_input_presentation(package: WorkbookPackage) -> None:
    """Kennzeichnet Eingaben und Anzeigen einheitlich; Zahlenformate bleiben erhalten."""
    styles = ET.fromstring(package.parts["xl/styles.xml"])
    formats = styles.find(tag("cellXfs"))
    assert formats is not None
    variants: dict[tuple[int, bool], str] = {}

    def style_index(original: str, calculated: bool) -> str:
        key = int(original), calculated
        if key not in variants:
            style = deepcopy(formats[key[0]])
            style.set("fillId", "8" if calculated else "0")
            style.set("borderId", "3")
            style.set("fontId", "2" if calculated else "6")
            for attribute in ("applyFill", "applyBorder", "applyFont"):
                style.set(attribute, "1")
            variants[key] = str(len(formats))
            formats.append(style)
        return variants[key]

    roots = {title: package.read_sheet(title) for title in package.sheets}
    for title, root in roots.items():
        if title in ("Hilfsdaten", "Klienten (alt)"):
            continue
        path = package.sheets[title]
        rel_path = path.replace("/worksheets/", "/worksheets/_rels/") + ".rels"
        if rel_path not in package.parts:
            continue
        for rel in ET.fromstring(package.parts[rel_path]):
            target = rel.get("Target", "")
            if "/tables/" not in target:
                continue
            table = ET.fromstring(package.parts["xl/tables/" + target.rsplit("/", 1)[-1]])
            left, top, right, bottom = range_boundaries(table.get("ref", ""))
            assert left and top and right and bottom
            columns = table.findall("m:tableColumns/m:tableColumn", NS)
            cells = cell_map(root)
            for number in range(left, right + 1):
                field = columns[number - left].get("name", "")
                first = cells.get(f"{get_column_letter(number)}{top + 1}")
                calculated_column = (
                    columns[number - left].find(tag("calculatedColumnFormula")) is not None
                    or first is not None
                    and first.find(tag("f")) is not None
                    or title == "Prüfkatalog"
                    and field not in ("Kundenkommentar", "Bearbeitung")
                    or title == "Fehlerliste"
                )
                default = style_index(first.get("s", "25") if first is not None else "25", calculated_column)
                cols = root.find(tag("cols"))
                if cols is not None:
                    for col in cols:
                        if int(col.get("min", "0")) == number == int(col.get("max", "0")):
                            col.set("style", default)
                for row in range(top + 1, bottom + 1):
                    ref = f"{get_column_letter(number)}{row}"
                    cell = cells.get(ref)
                    if cell is None:
                        cell = set_cell(root, ref, "", default)
                    calculated = calculated_column or cell.find(tag("f")) is not None
                    cell.set("s", style_index(cell.get("s", "25"), calculated))
        package.save_sheet(title, root)

    root = package.read_sheet("Prüfungen")
    set_cell(root, "F6", "Fehlerliste", "21")
    for row in range(7, 66):
        ref = f"F{row}"
        cells = cell_map(root)
        cell = cells.get(ref)
        if cell is not None and cell.find(tag("f")) is not None:
            formula = cell.find(tag("f"))
            assert formula is not None
            formula.text = cell.findtext(tag("f"), "").replace("Datensätze öffnen", "Fehlerliste öffnen")
        for column in "ABCDEFG":
            cell = cells.get(f"{column}{row}")
            if cell is not None:
                cell.set("s", style_index(cell.get("s", "25"), True))
    for col in root.findall("m:cols/m:col", NS):
        if int(col.get("min", "0")) >= 8:
            col.set("hidden", "1")
    cols = root.find(tag("cols"))
    assert cols is not None
    if not any(col.get("min") == "6" for col in cols):
        ET.SubElement(cols, tag("col"), {"min": "6", "max": "6", "width": "24", "customWidth": "1"})
        cols[:] = sorted(cols, key=lambda col: int(col.get("min", "0")))
    package.save_sheet("Prüfungen", root)

    for title, prefix in (("Bezugspersonen", "BP"), ("Kind-Bezugsperson", "KB"), ("Platzierungen", "PL")):
        root = package.read_sheet(title)
        cells = cell_map(root)
        text = package.value(cells.get("A1"))
        convention = f" Kennungen: {prefix}0001, {prefix}0002 usw.; eindeutig, dauerhaft, nicht wiederverwenden. Bestehende Kennungen behalten."
        if convention not in text:
            set_cell(root, "A1", text + convention, cells["A1"].get("s", "100"))
        row = root.find('m:sheetData/m:row[@r="1"]', NS)
        assert row is not None
        row.set("ht", "76")
        package.save_sheet(title, root)
    formats.set("count", str(len(formats)))
    package.parts["xl/styles.xml"] = serialize(styles, package.parts["xl/styles.xml"])

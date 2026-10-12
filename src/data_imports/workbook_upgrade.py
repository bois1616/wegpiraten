"""Bereitet eine Kundenkopie vor und erhält deren manuell gepflegte Eingaben."""

import hashlib
import re
from pathlib import Path
from xml.etree import ElementTree as ET

from loguru import logger
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import column_index_from_string, get_column_letter, range_boundaries
from pydantic import BaseModel

from data_imports.normalized_masterdata import TABLES
from data_imports.workbook_package import NS, WorkbookPackage, cell_map, serialize, set_cell, tag
from shared_modules.config import Config

_DIAGNOSTICS = {
    "Kinder": ("person", "Y", "Z", "AA", "AB", "AC"),
    "Aufträge": ("mandate", "AL", "AM", "AN", "AO", "AP"),
    "Betreuungen": ("mandate_person", "Z", "AA", "AB", "AC", "AD"),
    "Zuordnung MA": ("relation_mandate_emp", "K", "L", "M", "N", "O"),
    "Ansprechpersonen": ("masterdata_contact_person", "J", "K", "L", "M", "N"),
    "Berichte": ("report", "P", "Q", "R", "S", "T"),
}
_TITLES = {
    "assignment_history": "Zuständigkeitsverlauf",
    "related_person": "Bezugspersonen",
    "person_related_person": "Kind-Bezugsperson",
    "placement": "Platzierungen",
}
_LABELS = {
    "assignment_id": "Zeitraum-Nr",
    "mandate_id": "Auftrag-Nr",
    "employee_id": "Mitarbeiter-Nr",
    "role": "Rolle",
    "valid_from": "Gültig ab",
    "valid_to": "Gültig bis",
    "notes": "Bemerkungen",
    "related_person_id": "Bezugsperson-Nr",
    "first_name": "Vorname",
    "last_name": "Nachname",
    "street": "Strasse",
    "zip_code": "PLZ",
    "city": "Ort",
    "phone": "Telefon",
    "email": "E-Mail",
    "relationship_id": "Zuordnung-Nr",
    "person_id": "Kind-Nr",
    "is_legal_guardian": "Sorgeberechtigt?",
    "placement_id": "Platzierung-Nr",
    "placement_type": "Art der Platzierung",
    "place_name": "Einrichtung / Pflegefamilie",
    "start_date": "Beginn (falls bekannt)",
    "end_date": "Ende (falls bekannt)",
    "residence_street": "Wohnadresse Kind",
    "residence_zip_code": "PLZ Kind",
    "residence_city": "Wohnort Kind",
    "residence_phone": "Telefon Kind",
    "billing_email": "E-Mail Rechnung",
    "person_name": "Kind",
    "related_person_name": "Bezugsperson",
    "employee_name": "Mitarbeiter",
    "index_person_name": "Indexkind",
    "issue_err": "Fehler (berechnet)",
    "issue_hint": "Hinweis (berechnet)",
    "issue_kind": "Art (berechnet)",
    "issue_text": "Befund (berechnet)",
    "issue_run": "Zähler (berechnet)",
}


class UpgradeReport(BaseModel):
    """Prüfbarer Nachweis einer technischen Umstellung."""

    source_sha256: str
    output_sha256: str = ""
    preserved_input_cells: int = 0
    repaired_formula_cells: int = 0
    added_sheets: list[str] = []
    changed_parts: list[str] = []


def _inputs(package: WorkbookPackage, config: Config) -> dict[str, str]:
    """Vergleicht Rohwerte; berechnete Fremdschlüssel gehören nicht zu Eingaben."""
    values: dict[str, str] = {}
    for table_name, (_, entity) in TABLES.items():
        path, table = package.table(table_name)
        fields = [column.get("name", "") for column in table.findall("m:tableColumns/m:tableColumn", NS)]
        title = next(title for title in package.sheets if table_name in _sheet_table_names(package, title))
        cells = cell_map(package.read_sheet(title))
        left, top, _, bottom = range_boundaries(table.get("ref", ""))
        assert left and top and bottom
        for field in config.models[entity].fields:
            name = field.excel_column
            if (
                name not in fields
                or name in ("predecessor_mandate_id",)
                or table_name == "report"
                and name == "mandate_id"
            ):
                continue
            column = get_column_letter(left + fields.index(name))
            for number in range(top + 1, bottom + 1):
                cell = cells.get(f"{column}{number}")
                if cell is not None and cell.find(tag("f")) is None and package.value(cell):
                    values[f"{table_name}/{name}/{number}"] = package.value(cell)
    return values


def _sheet_table_names(package: WorkbookPackage, title: str) -> set[str]:
    """Löst Tabellen über die echten Beziehungen statt über Blattnamen auf."""
    sheet_path = package.sheets[title]
    rel_path = sheet_path.replace("/worksheets/", "/worksheets/_rels/") + ".rels"
    if rel_path not in package.parts:
        return set()
    names = set()
    for rel in ET.fromstring(package.parts[rel_path]):
        target = rel.get("Target", "")
        if "/tables/" in target:
            path = "xl/tables/" + target.rsplit("/", 1)[-1]
            names.add(ET.fromstring(package.parts[path]).get("displayName", ""))
    return names


def _extend_columns(package: WorkbookPackage, title: str, table_name: str, fields: list[str]) -> None:
    """Hängt Eingaben rechts an; bestehende Zellbezüge werden nicht verschoben."""
    root = package.read_sheet(title)
    path, table = package.table(table_name)
    columns = table.find(tag("tableColumns"))
    assert columns is not None
    left, top, right, bottom = range_boundaries(table.get("ref", ""))
    assert left and top and right and bottom
    cols = root.find(tag("cols"))
    assert cols is not None
    for field in fields:
        right += 1
        ET.SubElement(columns, tag("tableColumn"), {"id": str(len(columns) + 1), "name": field})
        set_cell(root, f"{get_column_letter(right)}{top - 1}", _LABELS[field], "21")
        set_cell(root, f"{get_column_letter(right)}{top}", field, "5")
        set_cell(root, f"{get_column_letter(right)}{top + 1}", "")
        ET.SubElement(cols, tag("col"), {"min": str(right), "max": str(right), "width": "28", "customWidth": "1"})
    columns.set("count", str(len(columns)))
    ref = f"{get_column_letter(left)}{top}:{get_column_letter(right)}{bottom}"
    table.set("ref", ref)
    autofilter = table.find(tag("autoFilter"))
    if autofilter is not None:
        autofilter.set("ref", ref)
    dimension = root.find(tag("dimension"))
    if dimension is not None:
        dimension.set("ref", f"A1:{get_column_letter(right)}2000")
    package.put_table(path, table)
    package.save_sheet(title, root)


def _base_sheet(title: str, fields: list[str], help_text: str) -> ET.Element:
    """Neue Eingabeblätter folgen dem bestehenden Drei-Zeilen-Kopf."""
    root = ET.Element(tag("worksheet"))
    ET.SubElement(root, tag("dimension"), {"ref": f"A1:{get_column_letter(len(fields))}203"})
    views = ET.SubElement(root, tag("sheetViews"))
    view = ET.SubElement(views, tag("sheetView"), {"workbookViewId": "0", "showGridLines": "0"})
    ET.SubElement(
        view,
        tag("pane"),
        {"xSplit": "2", "ySplit": "3", "topLeftCell": "C4", "activePane": "bottomRight", "state": "frozen"},
    )
    ET.SubElement(root, tag("sheetFormatPr"), {"defaultRowHeight": "18"})
    cols = ET.SubElement(root, tag("cols"))
    for number, field in enumerate(fields, 1):
        ET.SubElement(
            cols,
            tag("col"),
            {
                "min": str(number),
                "max": str(number),
                "width": "60" if field == "notes" else "26",
                "customWidth": "1",
                **({"hidden": "1"} if field.startswith("issue_") else {}),
            },
        )
    ET.SubElement(root, tag("sheetData"))
    set_cell(root, "A1", title, "20")
    set_cell(root, "B1", help_text, "26")
    for number, field in enumerate(fields, 1):
        set_cell(root, f"{get_column_letter(number)}2", _LABELS.get(field, field), "21")
        set_cell(root, f"{get_column_letter(number)}3", field, "5")
        set_cell(
            root,
            f"{get_column_letter(number)}4",
            "",
            "25" if field in ("start_date", "end_date", "valid_from", "valid_to") else "24",
        )
    data = root.find(tag("sheetData"))
    assert data is not None
    for row in data:
        if row.get("r") == "3":
            row.set("hidden", "1")
        if row.get("r") == "1":
            row.set("ht", "42")
            row.set("customHeight", "1")
    return root


def _validation(root: ET.Element, column: str, formula: str) -> None:
    """Reguläre Listenauswahl mit klarer Fehlermeldung."""
    validations = root.find(tag("dataValidations"))
    if validations is None:
        validations = ET.SubElement(root, tag("dataValidations"), {"count": "0"})
    node = ET.SubElement(
        validations,
        tag("dataValidation"),
        {
            "type": "list",
            "allowBlank": "1",
            "showErrorMessage": "1",
            "errorStyle": "stop",
            "errorTitle": "Auswahl prüfen",
            "error": "Bitte einen Wert aus der Liste auswählen.",
            "sqref": f"{column}4:{column}203",
        },
    )
    ET.SubElement(node, tag("formula1")).text = formula
    validations.set("count", str(len(validations)))


def _repair(package: WorkbookPackage, title: str, definition: tuple[str, ...], report: UpgradeReport) -> None:
    """Füllt verlorene Formeln und ersetzt fragile Vorgängerzähler."""
    table_name, err, hint, kind, text, run = definition
    root = package.read_sheet(title)
    cells = cell_map(root)
    path, table = package.table(table_name)
    _, top, _, bottom = range_boundaries(table.get("ref", ""))
    assert top and bottom
    prototype: dict[str, str] = {}
    for name, cell in cells.items():
        if name.endswith("6") and re.fullmatch(r"[A-Z]+6", name):
            formula = cell.find(tag("f"))
            if formula is not None and formula.text and "#REF!" not in formula.text:
                prototype[re.findall(r"^[A-Z]+", name)[0]] = "=" + formula.text
    # Bekannte drei verlorene Bericht-FKs werden aus der sichtbaren Auswahl berechnet.
    if title == "Berichte":
        prototype["C"] = '=IF($B6="","",IF(ISNUMBER(FIND(" ",$B6)),LEFT($B6,FIND(" ",$B6)-1),$B6))'
        prototype[err] = prototype[err].replace(
            '&IF(AND($C6<>"",', '&IF($C6=""," · Auftragsverweis fehlt","")&IF(AND($C6<>"",'
        )
    if title == "Aufträge":
        prototype[err] = prototype[err].replace(
            ",4,900))", '&IF($E6=""," · Kostenträger fehlt","")&IF($C6=""," · Leistungsbesteller fehlt",""),4,900))'
        )
    if title == "Betreuungen":
        scoped = (
            'COUNTIF(reportable_payers,IFERROR(INDEX(Aufträge!$E$4:$E$1999,MATCH($A6,Aufträge!$A$4:$A$1999,0)),""))>0'
        )
        prototype[hint] = prototype[hint].replace('AND($D6<>"",$F6="")', f'AND($D6<>"",$E6="Nein",$F6="",{scoped})')
        prototype[hint] = prototype[hint].replace(
            ",4,900))",
            f'&IF(AND($D6<>"",NOT(OR($E6="Ja",$E6="Nein")),{scoped})," · Austritt geplant: Ja/Nein fehlt","")&IF(AND($D6<>"",$E6="Nein",$F6="Anderer",$G6="",{scoped})," · Anderer Austrittsgrund ohne Freitext",""),4,900))',
        )
    max_row = max(int(re.findall(r"\d+$", name)[0]) for name in cells)
    for column, formula in prototype.items():
        if column == run:
            continue
        for number in range(top + 1, max_row + 1):
            name = f"{column}{number}"
            cell = cells.get(name)
            old = cell.find(tag("f")) if cell is not None else None
            if cell is None or not package.value(cell) or old is not None:
                updated = Translator(formula, origin=f"{column}6").translate_formula(name)
                if old is None or "=" + (old.text or "") != updated:
                    set_cell(root, name, updated, cell.get("s", "26") if cell is not None else "26", True)
                    report.repaired_formula_cells += 1
    for number in range(top + 1, max_row + 1):
        formula = f'COUNTIF(INDEX(${kind}:${kind},ROW(${kind}${top})+1):{kind}{number},"?*")'
        set_cell(root, f"{run}{number}", formula, "26", True)
        report.repaired_formula_cells += 1
    columns = table.find(tag("tableColumns"))
    assert columns is not None
    for index, column in enumerate(columns, 1):
        letter = get_column_letter(index)
        if letter in prototype:
            formula = (
                f'COUNTIF(INDEX(${kind}:${kind},ROW(${kind}${top})+1):{kind}{top + 1},"?*")'
                if letter == run
                else Translator(prototype[letter], origin=f"{letter}6")
                .translate_formula(f"{letter}{top + 1}")
                .lstrip("=")
            )
            node = column.find(tag("calculatedColumnFormula"))
            if node is None:
                node = ET.SubElement(column, tag("calculatedColumnFormula"))
            node.text = formula
    package.put_table(path, table)
    package.save_sheet(title, root)


def _add_input_sheets(package: WorkbookPackage, config: Config) -> None:
    """Leere Zusatzlisten mit Eingabeprüfung und lesbaren Namensverweisen."""
    helpers = {
        "assignment_history": ["index_person_name", "employee_name"],
        "related_person": [],
        "person_related_person": ["person_name", "related_person_name"],
        "placement": ["person_name"],
    }
    for table, title in _TITLES.items():
        fields = [field.excel_column for field in config.models[table].fields if field.excel_column]
        original = list(fields)
        fields += helpers[table] + ["issue_err", "issue_hint", "issue_kind", "issue_text", "issue_run"]
        root = _base_sheet(
            title,
            fields,
            "Nur Eingaben pflegen. Berechnete Spalten nicht überschreiben. Bekannte Angaben erfassen; keine Historie erfinden.",
        )
        positions = {field: get_column_letter(index) for index, field in enumerate(fields, 1)}
        for number in range(4, 204):
            key = f"$A{number}"
            last_input = get_column_letter(len(original))
            problems = [f'IF(COUNTIF({table}[{original[0]}],{key})>1," · Nummer doppelt","")']
            hints: list[str] = []
            for field, target, target_field in (
                ("person_id", "person", "person_id"),
                ("related_person_id", "related_person", "related_person_id"),
                ("mandate_id", "mandate", "mandate_id"),
                ("employee_id", "masterdata_employee", "employee_id"),
            ):
                if field not in positions or field == original[0]:
                    continue
                value = f"${positions[field]}{number}"
                problems.append(
                    f'IF(OR({value}="",COUNTIF({target}[{target_field}],{value})=0)," · {_LABELS[field]} fehlt oder ist unbekannt","")'
                )
            if table == "person_related_person":
                role = f"${positions['role']}{number}"
                guardian = f"${positions['is_legal_guardian']}{number}"
                problems.append(
                    f'IF(COUNTIF(relationship_roles,{role})=0," · Beziehungsrolle fehlt oder ist unbekannt","")'
                )
                hints.append(
                    f'IF(AND({guardian}<>"",COUNTIF(guardian_values,{guardian})=0)," · Sorgeberechtigung: Ja/Nein/Unbekannt","")'
                )
            if table == "assignment_history":
                start, end, role, mandate = (
                    f"${positions[field]}{number}" for field in ("valid_from", "valid_to", "role", "mandate_id")
                )
                problems += [
                    f'IF(NOT(ISNUMBER({start}))," · Gültig-ab fehlt oder ist kein Datum","")',
                    f'IF(NOT(OR({role}="P",{role}="S"))," · Rolle P/S fehlt","")',
                ]
                hints.append(
                    f'IF(AND({role}="P",ISNUMBER({start}),COUNTIFS(assignment_history[mandate_id],{mandate},assignment_history[role],"P",assignment_history[valid_from],"<="&IF({end}="",DATE(9999,12,31),{end}),assignment_history[valid_to],">="&{start})+COUNTIFS(assignment_history[mandate_id],{mandate},assignment_history[role],"P",assignment_history[valid_from],"<="&IF({end}="",DATE(9999,12,31),{end}),assignment_history[valid_to],"")>1)," · P-Zeiträume überschneiden sich","")'
                )
            if table in ("placement", "assignment_history"):
                first, last = (
                    ("valid_from", "valid_to") if table == "assignment_history" else ("start_date", "end_date")
                )
                start, end = (f"${positions[field]}{number}" for field in (first, last))
                hints += [
                    f'IF(AND(ISNUMBER({start}),ISNUMBER({end}),{end}<{start})," · Ende vor Beginn","")',
                    f'IF(OR(AND({start}<>"",NOT(ISNUMBER({start}))),AND({end}<>"",NOT(ISNUMBER({end}))))," · Datum als Excel-Datum erfassen","")',
                ]
            if "email" in positions:
                email = f"${positions['email']}{number}"
                hints.append(
                    f'IF(AND({email}<>"",OR(ISERROR(SEARCH("@",{email})),ISERROR(SEARCH(".",{email})),ISNUMBER(SEARCH(" ",{email}))))," · E-Mail prüfen","")'
                )
            err, hint, kind, text, run = (
                positions[field] for field in ("issue_err", "issue_hint", "issue_kind", "issue_text", "issue_run")
            )
            set_cell(
                root,
                f"{err}{number}",
                f'IF(AND({key}="",COUNTIF(A{number}:{last_input}{number},"?*")+COUNT(A{number}:{last_input}{number})>0),"Nummer fehlt",IF({key}="","",MID(""&'
                + "&".join(problems)
                + ",4,900)))",
                "26",
                True,
            )
            set_cell(
                root,
                f"{hint}{number}",
                f'IF({key}="","",MID(""' + "".join("&" + value for value in hints) + ",4,900))",
                "26",
                True,
            )
            set_cell(
                root,
                f"{kind}{number}",
                f'IF({err}{number}<>"","Fehler",IF({hint}{number}<>"","Hinweis",""))',
                "26",
                True,
            )
            set_cell(
                root,
                f"{text}{number}",
                f'{err}{number}&IF(AND({err}{number}<>"",{hint}{number}<>"")," · ","")&{hint}{number}',
                "26",
                True,
            )
            set_cell(root, f"{run}{number}", f'COUNTIF(${kind}$4:{kind}{number},"?*")', "26", True)
            for helper, value_field, source_ids, source_values in (
                ("person_name", "person_id", "Kinder!$A$4:$A$2000", "Kinder!$W$4:$W$2000"),
                (
                    "related_person_name",
                    "related_person_id",
                    "Bezugspersonen!$A$4:$A$203",
                    "Bezugspersonen!$C$4:$C$203",
                ),
                ("employee_name", "employee_id", "Mitarbeiter!$B$3:$B$200", "Mitarbeiter!$C$3:$C$200"),
                ("index_person_name", "mandate_id", "Aufträge!$A$4:$A$1999", "Aufträge!$V$4:$V$1999"),
            ):
                if helper in positions:
                    key_field = f"${positions[value_field]}{number}"
                    set_cell(
                        root,
                        f"{positions[helper]}{number}",
                        f'IF({key_field}="","",IFERROR(INDEX({source_values},MATCH({key_field},{source_ids},0))&"","unbekannt"))',
                        "26",
                        True,
                    )
        for field, name in (
            ("mandate_id", "upgrade_mandate_ids"),
            ("employee_id", "upgrade_employee_ids"),
            ("person_id", "upgrade_person_ids"),
            ("related_person_id", "upgrade_related_person_ids"),
        ):
            if field in positions and field != original[0]:
                _validation(root, positions[field], name)
        if "role" in positions:
            _validation(root, positions["role"], '"P,S"' if table == "assignment_history" else "relationship_roles")
        if "is_legal_guardian" in positions:
            _validation(root, positions["is_legal_guardian"], "guardian_values")
        package.add_sheet(title, root, table, fields)
        path, definition = package.table(table)
        columns = definition.find(tag("tableColumns"))
        assert columns is not None
        first_row = cell_map(root)
        for number, column in enumerate(columns, 1):
            formula = first_row[f"{get_column_letter(number)}4"].find(tag("f"))
            if formula is not None:
                ET.SubElement(column, tag("calculatedColumnFormula")).text = formula.text
        package.put_table(path, definition)
    for name, reference in {
        "upgrade_mandate_ids": "Aufträge!$A$4:$A$1999",
        "upgrade_employee_ids": "Mitarbeiter!$B$3:$B$200",
        "upgrade_person_ids": "Kinder!$A$4:$A$2000",
        "upgrade_related_person_ids": "Bezugspersonen!$A$4:$A$203",
    }.items():
        package.define_name(name, reference)


class CheckSpec(BaseModel):
    """Feste Prüfung, getrennt von der dynamischen offenen Ansicht."""

    check_id: str
    kind: str
    sheet: str
    title: str
    explanation: str
    formula: str
    comment: str = ""
    disposition: str = ""


def _catalog(package: WorkbookPackage, config: Config) -> None:
    """Bindet Kommentare an feste Prüfungskennungen und erweitert fehlende Kontrollen."""
    old = cell_map(package.read_sheet("Prüfungen"))
    comments = {
        package.value(old.get(f"C{row}")): package.value(old.get(f"G{row}"))
        for row in range(7, 50)
        if package.value(old.get(f"G{row}"))
    }
    checks: list[CheckSpec] = []
    for row in range(7, 50):
        title = package.value(old.get(f"J{row}"))
        formula = old[f"L{row}"].findtext(tag("f"), "")
        if row == 39:
            formula = 'SUMPRODUCT((Betreuungen!$D$4:$D$1998<>"")*(Betreuungen!$E$4:$E$1998="Nein")*(Betreuungen!$F$4:$F$1998="")*(COUNTIF(reportable_payers,IFERROR(INDEX(Aufträge!$E$4:$E$1999,MATCH(Betreuungen!$A$4:$A$1998,Aufträge!$A$4:$A$1999,0)),""))>0))'
        if row == 47:
            formula = "SUMPRODUCT((Aufträge!$A$4:$A$1999<>\"\")*(COUNTIFS('Zuordnung MA'!$A$4:$A$2000,Aufträge!$A$4:$A$1999,'Zuordnung MA'!$C$4:$C$2000,\"P\")=0))"
        comment = comments.get(title, "")
        disposition = ""
        if comment:
            disposition = "Betroffene Datensätze in der Fehlerliste öffnen; fachliche Korrektur durch den Kunden."
            if row == 39:
                disposition = "Technische Prüfung korrigiert: Grund nur bei Austritt geplant = Nein und meldepflichtigem Kostenträger."
            if row == 49:
                disposition = (
                    "Bestätigung erhalten. Meldepflichtige Betreuungen ohne Mapping bleiben zur Prüfung sichtbar."
                )
        checks.append(
            CheckSpec(
                check_id=f"P{row - 6:03d}",
                kind=package.value(old.get(f"H{row}")),
                sheet=package.value(old.get(f"I{row}")),
                title=title,
                explanation=package.value(old.get(f"K{row}")),
                formula=formula,
                comment=comment,
                disposition=disposition,
            )
        )
    additions = [
        (
            "Fehler",
            "Berichte",
            "Bericht-Nr doppelt vergeben",
            'SUMPRODUCT((Berichte!$A$4:$A$1990<>"")*(COUNTIF(Berichte!$A$4:$A$1990,Berichte!$A$4:$A$1990)>1))',
        ),
        (
            "Fehler",
            "Berichte",
            "Auftragsverweis fehlt oder ist unbekannt",
            'SUMPRODUCT((Berichte!$A$4:$A$1990<>"")*((Berichte!$C$4:$C$1990="")+(COUNTIF(Aufträge!$A$4:$A$1999,Berichte!$C$4:$C$1990)=0)>0))',
        ),
        ("Hinweis", "Berichte", "Bericht überfällig", 'COUNTIF(Berichte!$O$4:$O$1990,"überfällig")'),
        ("Hinweis", "Berichte", "Erledigt-Datum fehlt", 'COUNTIF(Berichte!$O$4:$O$1990,"Datum fehlt")'),
        (
            "Hinweis",
            "Berichte",
            "Fälligkeit fehlt",
            'SUMPRODUCT((Berichte!$A$4:$A$1990<>"")*(Berichte!$E$4:$E$1990=""))',
        ),
        (
            "Fehler",
            "Aufträge",
            "Kostenträger fehlt",
            'SUMPRODUCT((Aufträge!$A$4:$A$1999<>"")*(Aufträge!$E$4:$E$1999=""))',
        ),
        (
            "Fehler",
            "Aufträge",
            "Leistungsbesteller fehlt",
            'SUMPRODUCT((Aufträge!$A$4:$A$1999<>"")*(Aufträge!$C$4:$C$1999=""))',
        ),
        (
            "Hinweis",
            "Berichte",
            "Zuständigkeit am Fälligkeitstag fehlt",
            'SUMPRODUCT((Berichte!$A$4:$A$1990<>"")*(Berichte!$C$4:$C$1990<>"")*(Berichte!$M$4:$M$1990=""))',
        ),
    ]
    for table, title in _TITLES.items():
        additions += [
            ("Fehler", title, "Eingaben / Verweise prüfen", f'COUNTIF({table}[issue_err],"?*")'),
            ("Hinweis", title, "Zusatzdaten prüfen", f'COUNTIF({table}[issue_hint],"?*")'),
        ]
    for kind, sheet, title, formula in additions:
        checks.append(
            CheckSpec(
                check_id=f"P{len(checks) + 1:03d}",
                kind=kind,
                sheet=sheet,
                title=title,
                explanation="Die Fehlerliste nennt Datensatz und Zeile. Änderungen in der jeweiligen Eingabetabelle vornehmen.",
                formula=formula,
            )
        )
    fields = [
        "Prüfung-ID",
        "Art",
        "Blatt",
        "Prüfung",
        "Erläuterung",
        "Anzahl",
        "Kundenkommentar",
        "Bearbeitung",
        "Status",
        "Rang",
    ]
    root = _base_sheet(
        "Prüfkatalog",
        fields,
        "Kommentare hier bei der festen Prüfung pflegen. Die offene Ansicht zieht diese automatisch nach.",
    )
    for row, check in enumerate(checks, 4):
        for column, value in enumerate(
            (
                check.check_id,
                check.kind,
                check.sheet,
                check.title,
                check.explanation,
                "",
                check.comment,
                check.disposition,
            ),
            1,
        ):
            set_cell(root, f"{get_column_letter(column)}{row}", value, "26")
        set_cell(root, f"F{row}", check.formula, "26", True)
        set_cell(root, f"I{row}", f'IF(F{row}>0,"offen","ok")', "26", True)
        set_cell(root, f"J{row}", f'IF(F{row}>0,COUNTIF($F$4:F{row},">0"),0)', "26", True)
    # Kleine Wertelisten sind sichtbar im Katalog und werden nicht aus Kontaktdaten abgeleitet.
    for index, value in enumerate(config.reporting.accordix_payer_ids, 4):
        set_cell(root, f"M{index}", value)
    set_cell(root, "M2", "Meldepflichtige Kostenträger", "21")
    roles = ("Mutter", "Vater", "Beistand", "Pflegeperson", "Bezugsperson", "Sonstige")
    for index, value in enumerate(roles, 4):
        set_cell(root, f"N{index}", value)
    set_cell(root, "N2", "Beziehungsrollen", "21")
    for index, value in enumerate(("Ja", "Nein", "Unbekannt"), 4):
        set_cell(root, f"O{index}", value)
    set_cell(root, "O2", "Sorgeberechtigung", "21")
    package.add_sheet("Prüfkatalog", root, "check_catalog", fields, len(checks) + 3)
    package.define_name(
        "reportable_payers", f"Prüfkatalog!$M$4:$M${max(4, len(config.reporting.accordix_payer_ids) + 3)}"
    )
    package.define_name("relationship_roles", "Prüfkatalog!$N$4:$N$9")
    package.define_name("guardian_values", "Prüfkatalog!$O$4:$O$6")
    view = package.read_sheet("Prüfungen")
    set_cell(
        view,
        "A2",
        "Offene Prüfungen. Kommentare gehören in den festen Prüfkatalog; die Fehlerliste zeigt Datensätze und verlinkt ihre Zeilen.",
        "26",
    )
    set_cell(view, "B4", 'COUNTIFS(check_catalog[Art],"Fehler",check_catalog[Anzahl],">0")', "26", True)
    set_cell(view, "C4", f'COUNTIF(check_catalog[Anzahl],">0")&" von {len(checks)} Prüfungen sind offen."', "26", True)
    set_cell(view, "F6", "Einträge", "21")
    set_cell(view, "G6", "Kundenkommentar", "21")
    for row in range(7, len(checks) + 7):
        for column, field in enumerate(("Art", "Blatt", "Prüfung", "Erläuterung", "Anzahl"), 1):
            set_cell(
                view,
                f"{get_column_letter(column)}{row}",
                f"IFERROR(INDEX(check_catalog[{field}],MATCH(ROW()-6,check_catalog[Rang],0))"
                + ("" if field == "Anzahl" else '&""')
                + ',"")',
                "26",
                True,
            )
        set_cell(view, f"F{row}", f'IF(A{row}="","",HYPERLINK("#\'Fehlerliste\'!A4","Datensätze öffnen"))', "26", True)
        set_cell(
            view,
            f"G{row}",
            'IFERROR(INDEX(check_catalog[Kundenkommentar],MATCH(ROW()-6,check_catalog[Rang],0))&"","")',
            "26",
            True,
        )
    package.save_sheet("Prüfungen", view)


def _due_assignment(package: WorkbookPackage) -> None:
    """Auch die sichtbare Berichtsliste verwendet die Fälligkeit und nicht TODAY."""
    root = package.read_sheet("Berichte")
    cells = cell_map(root)
    old = cells["M6"].findtext(tag("f"), "")
    for row in range(4, 1991):
        current = Translator("=" + old, origin="M6").translate_formula(f"M{row}").lstrip("=")
        count = f'COUNTIFS(assignment_history[mandate_id],$C{row},assignment_history[role],"P",assignment_history[valid_from],"<="&$E{row},assignment_history[valid_to],">="&$E{row})+COUNTIFS(assignment_history[mandate_id],$C{row},assignment_history[role],"P",assignment_history[valid_from],"<="&$E{row},assignment_history[valid_to],"")'
        formula = f'IF($C{row}="","",IF(COUNTIF(assignment_history[mandate_id],$C{row})=0,{current},IF(AND(ISNUMBER($E{row}),({count})=1),IFERROR(INDEX(assignment_history[employee_id],MATCH(1,INDEX((assignment_history[mandate_id]=$C{row})*(assignment_history[role]="P")*ISNUMBER(assignment_history[valid_from])*(assignment_history[valid_from]<=$E{row})*((assignment_history[valid_to]="")+(assignment_history[valid_to]>=$E{row})),0),0)),""),"")))'
        set_cell(root, f"M{row}", formula, "26", True)
    path, table = package.table("report")
    column = table.findall("m:tableColumns/m:tableColumn", NS)[12]
    node = column.find(tag("calculatedColumnFormula"))
    assert node is not None
    node.text = cell_map(root)["M4"].findtext(tag("f"), "")
    package.put_table(path, table)
    package.save_sheet("Berichte", root)


def _findings(package: WorkbookPackage) -> None:
    """Erweitert die vorhandene Liste um neue Blätter und direkte Zeilenlinks."""
    source = [(title, definition[4], definition[5], definition[3]) for title, definition in _DIAGNOSTICS.items()]
    for table, title in _TITLES.items():
        _, definition = package.table(table)
        fields = [column.get("name", "") for column in definition.findall("m:tableColumns/m:tableColumn", NS)]
        source.append(
            (
                title,
                get_column_letter(fields.index("issue_text") + 1),
                get_column_letter(fields.index("issue_run") + 1),
                get_column_letter(fields.index("issue_kind") + 1),
            )
        )
    root = package.read_sheet("Fehlerliste")
    for index, (title, _, run, _) in enumerate(source, 1):
        end = (
            203
            if title in _TITLES.values()
            else 1990
            if title == "Berichte"
            else 1998
            if title == "Betreuungen"
            else 1999
            if title == "Aufträge"
            else 2000
        )
        set_cell(root, f"K{index}", title, "26")
        set_cell(root, f"L{index}", f"MAX('{title}'!${run}$4:${run}${end})", "26", True)
        set_cell(root, f"M{index}", "0" if index == 1 else f"M{index - 1}+L{index - 1}", "26", True)
    total_row = len(source) + 1
    set_cell(root, f"L{total_row}", f"SUM(L1:L{len(source)})", "26", True)
    set_cell(
        root,
        "A2",
        f'L{total_row}&" Datensätze mit Befund. Zeilenlinks öffnen die Eingabe. Bei mehr als 600 Befunden werden die ersten 600 angezeigt."',
        "26",
        True,
    )
    for row in range(5, 605):
        set_cell(root, f"G{row}", f'IF(ROW()-4>$L${total_row},"",MATCH(ROW()-5,$M$1:$M${len(source)},1))', "26", True)
        set_cell(root, f"H{row}", f'IF($G{row}="","",ROW()-4-INDEX($M$1:$M${len(source)},$G{row}))', "26", True)
        match_cases, key_cases, kind_cases, text_cases = [], [], [], []
        for index, (title, text, run, kind) in enumerate(source, 1):
            end = (
                203
                if title in _TITLES.values()
                else 1990
                if title == "Berichte"
                else 1998
                if title == "Betreuungen"
                else 1999
                if title == "Aufträge"
                else 2000
            )
            match_cases.append((index, f"MATCH($H{row},'{title}'!${run}$4:${run}${end},0)"))
            key_cases.append((index, f"INDEX('{title}'!$A$4:$A${end},$I{row})&\"\""))
            kind_cases.append((index, f"INDEX('{title}'!${kind}$4:${kind}${end},$I{row})&\"\""))
            text_cases.append((index, f"INDEX('{title}'!${text}$4:${text}${end},$I{row})&\"\""))

        def choices(cases: list[tuple[int, str]]) -> str:
            expression = '""'
            for index, value in reversed(cases):
                expression = f"IF($G{row}={index},{value},{expression})"
            return expression

        set_cell(root, f"I{row}", f'IF($H{row}="","",{choices(match_cases)})', "26", True)
        set_cell(root, f"A{row}", f'IF($G{row}="","",INDEX($K$1:$K${len(source)},$G{row}))', "26", True)
        set_cell(
            root,
            f"B{row}",
            f'IF($I{row}="","",HYPERLINK("#\'"&A{row}&"\'!A"&($I{row}+3),"Zeile "&($I{row}+3)))',
            "26",
            True,
        )
        for column, cases in (("C", key_cases), ("D", kind_cases), ("E", text_cases)):
            set_cell(root, f"{column}{row}", f'IF($I{row}="","",{choices(cases)})', "26", True)
    path, definition = package.table("fehlerliste")
    definition.set("ref", "A4:E604")
    node = definition.find(tag("autoFilter"))
    if node is not None:
        node.set("ref", "A4:E604")
    package.put_table(path, definition)
    package.save_sheet("Fehlerliste", root)


def finish_presentation(package: WorkbookPackage) -> None:
    """Ergänzt lesbare Katalogtexte und Textformat für neue Telefon-/PLZ-Eingaben."""
    from copy import deepcopy

    styles = ET.fromstring(package.parts["xl/styles.xml"])
    formats = styles.find(tag("cellXfs"))
    assert formats is not None
    text_style = deepcopy(formats[24])
    text_style.set("numFmtId", "49")
    text_style.set("applyNumberFormat", "1")
    text_index = str(len(formats))
    formats.append(text_style)
    wrap_style = deepcopy(formats[26])
    wrap_style.set("applyAlignment", "1")
    alignment = wrap_style.find(tag("alignment"))
    if alignment is None:
        alignment = ET.SubElement(wrap_style, tag("alignment"))
    alignment.set("wrapText", "1")
    alignment.set("vertical", "top")
    wrap_index = str(len(formats))
    formats.append(wrap_style)
    formats.set("count", str(len(formats)))
    package.parts["xl/styles.xml"] = serialize(styles, package.parts["xl/styles.xml"])
    for title, columns in {
        "Kinder": ("AE", "AG"),
        "Ansprechpersonen": ("O",),
        "Bezugspersonen": ("E", "G"),
        "Platzierungen": ("F", "H"),
    }.items():
        root = package.read_sheet(title)
        for cell in root.findall(".//m:c", NS):
            reference = cell.get("r", "")
            if re.sub(r"\d", "", reference) in columns and int(re.sub(r"\D", "", reference)) >= 4:
                cell.set("s", text_index)
        cols = root.find(tag("cols"))
        assert cols is not None
        for column in columns:
            number = str(column_index_from_string(column))
            for node in cols:
                if node.get("min") == number and node.get("max") == number:
                    node.set("style", text_index)
        package.save_sheet(title, root)
    root = package.read_sheet("Prüfkatalog")
    cols = root.find(tag("cols"))
    assert cols is not None
    for node in cols:
        number = int(node.get("min", "0"))
        if number in (4, 5, 7, 8):
            node.set("width", "65" if number in (4, 5) else "48")
        if number in (9, 10):
            node.set("hidden", "1")
    for row in root.findall("m:sheetData/m:row", NS):
        if int(row.get("r", "0")) >= 4:
            row.set("ht", "84")
            row.set("customHeight", "1")
            for cell in row:
                if re.sub(r"\d", "", cell.get("r", "")) in ("D", "E", "G", "H"):
                    cell.set("s", wrap_index)
    package.save_sheet("Prüfkatalog", root)


def prepare_workbook(config: Config, source: Path, destination: Path) -> UpgradeReport:
    """Bereitet eine separate, rückverfolgbare Kopie des bestätigten Formats vor."""
    if source.resolve() == destination.resolve():
        raise ValueError("Die Originalmappe darf nicht überschrieben werden.")
    package = WorkbookPackage(source)
    original_parts = dict(package.parts)
    inputs = _inputs(package, config)
    report = UpgradeReport(source_sha256=hashlib.sha256(source.read_bytes()).hexdigest())
    for title, definition in _DIAGNOSTICS.items():
        _repair(package, title, definition, report)
    _extend_columns(
        package, "Kinder", "person", ["residence_street", "residence_zip_code", "residence_city", "residence_phone"]
    )
    _extend_columns(package, "Ansprechpersonen", "masterdata_contact_person", ["phone", "email"])
    _extend_columns(package, "Kostenträger", "masterdata_payer", ["billing_email"])
    _add_input_sheets(package, config)
    _catalog(package, config)
    _due_assignment(package)
    _findings(package)
    finish_presentation(package)
    if _inputs(package, config) != inputs:
        raise ValueError("Eingabewerte wurden verändert; Ausgabe wird nicht geschrieben.")
    report.preserved_input_cells = len(inputs)
    report.added_sheets = ["Prüfkatalog", *_TITLES.values()]
    package.write(destination)
    report.changed_parts = sorted(
        name for name in set(original_parts) | set(package.parts) if original_parts.get(name) != package.parts.get(name)
    )
    report.output_sha256 = hashlib.sha256(destination.read_bytes()).hexdigest()
    destination.with_suffix(".umstellung.json").write_text(report.model_dump_json(indent=2), encoding="utf-8")
    logger.info("Kundenkopie vorbereitet: {} ({} Eingabezellen unverändert)", destination, len(inputs))
    return report

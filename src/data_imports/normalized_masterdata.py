"""Importiert das neue Workbook in die wegwerfbare Arbeitsdatenbank."""

import re
import shutil
import sqlite3
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import pandas as pd
from loguru import logger
from openpyxl.utils.cell import range_boundaries
from pydantic import BaseModel, ValidationError, create_model

from data_imports.mandate_schema import create_mandate_view
from shared_modules.config import Config
from shared_modules.internal_client import (
    INTERNAL_ALLOWED_DIRECT_EFFORT,
    INTERNAL_ALLOWED_INDIRECT_EFFORT,
    INTERNAL_ALLOWED_TRAVEL_TIME,
)

# Physische Namen der bereits vorhandenen Kostenträger-/Bestellertabellen bleiben.
TABLES: dict[str, tuple[str, str]] = {
    "masterdata_employee": ("employees", "employee"),
    "masterdata_payer": ("payer", "payer"),
    "masterdata_service_requester": ("service_requester", "service_requester"),
    "masterdata_tenant": ("masterdata_tenant", "tenant"),
    "service_types": ("service_types", "service_type"),
    "masterdata_contact_person": ("masterdata_contact_person", "contact_person"),
    "person": ("person", "person"),
    "mandate": ("mandate", "mandate"),
    "mandate_person": ("mandate_person", "mandate_person"),
    "relation_mandate_emp": ("relation_mandate_emp", "mandate_employee_relation"),
    "report": ("report", "report"),
}

# Fremdverweise werden diagnostiziert, nicht durch DB-Constraints verworfen.
FOREIGN_KEYS: dict[str, list[tuple[str, str, str, int | None]]] = {
    "masterdata_contact_person": [("service_requester_id", "service_requester", "service_requester_id", 13)],
    "mandate": [
        ("service_type_id", "service_types", "service_type_id", 29),
        ("service_requester_id", "service_requester", "service_requester_id", 13),
        ("payer_id", "payer", "payer_id", None),
        ("tenant_id", "masterdata_tenant", "tenant_id", None),
        ("contact_person_id", "masterdata_contact_person", "contact_person_id", 13),
        ("predecessor_mandate_id", "mandate", "mandate_id", 9),
    ],
    "mandate_person": [("mandate_id", "mandate", "mandate_id", 4), ("person_id", "person", "person_id", 5)],
    "relation_mandate_emp": [("mandate_id", "mandate", "mandate_id", 12), ("employee_id", "employees", "emp_id", 12)],
    "report": [("mandate_id", "mandate", "mandate_id", 38)],
}


class MasterdataFinding(BaseModel):
    """Ein nachbearbeitbarer Datenbefund mit Tabellen- und Zeilenbezug."""

    invariant: int | None = None
    table: str
    key: str
    source_row: int | None = None
    message: str


def import_normalized(config: Config, source: Path) -> int:
    """Liest alle Tabellen vor dem Schreiben und erhält fachliche Datenbefunde."""
    from data_imports.import_masterdata import get_type_from_str, map_row, sql_type

    workbook = openpyxl.load_workbook(source, data_only=True)
    frames: dict[str, tuple[pd.DataFrame, int]] = {}
    try:
        for excel, (_, entity) in TABLES.items():
            fields = config.models[entity].fields
            sheet = next((ws for ws in workbook if excel in ws.tables), None)
            if sheet is None:
                raise ValueError(f"Erforderliche Excel-Tabelle fehlt: {excel}")
            left, top, right, bottom = range_boundaries(sheet.tables[excel].ref)
            if left is None or top is None or right is None or bottom is None:
                raise ValueError(f"{excel}: Tabellenbereich ist ungültig")
            values = list(sheet.iter_rows(min_row=top, max_row=bottom, min_col=left, max_col=right, values_only=True))
            frame = pd.DataFrame(values[1:], columns=values[0])
            missing = [f.excel_column for f in fields if f.excel_column not in frame.columns]
            if missing:
                raise ValueError(f"{excel}: erforderliche Spalten fehlen: {missing}")
            frames[excel] = (frame, top + 1)
    finally:
        workbook.close()

    db = config.get_db_path()
    db.parent.mkdir(parents=True, exist_ok=True)
    if db.exists():
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        with (
            sqlite3.connect(db) as original,
            sqlite3.connect(db.with_name(f"{db.stem}_vor_import_{stamp}.sqlite3")) as backup,
        ):
            original.backup(backup)
    findings: list[MasterdataFinding] = []

    def finding(table: str, key: str, message: str, row: int | None = None, invariant: int | None = None) -> None:
        findings.append(MasterdataFinding(table=table, key=key, message=message, source_row=row, invariant=invariant))
        logger.warning("{} {} Zeile {}: {}", table, key, row, message)

    total = 0
    with sqlite3.connect(db) as conn:
        # Der Monatslauf beginnt mit einer gelöschten DB. Ein versehentlicher Wiederimport
        # in eine bestehende DB darf keine alten Leistungsdaten an neue Stammdaten hängen.
        existing = {r[0] for r in conn.execute("SELECT name FROM sqlite_master WHERE type='table'")}
        if existing:
            raise ValueError("Arbeitsdatenbank ist nicht leer. Nach Sicherung löschen und den Lauf neu beginnen.")
        for excel, (target, entity) in TABLES.items():
            fields = config.models[entity].fields
            keys = [f.name for f in fields if f.primary_key]
            columns = [f"{f.name} {sql_type(f.type)}" for f in fields]
            columns += ["source_row INTEGER", "is_active INTEGER NOT NULL DEFAULT 1"]
            if target == "person":
                columns.append("family_key TEXT")
            if target == "mandate":
                columns.append("needs_review INTEGER NOT NULL DEFAULT 0")
            conn.execute(f"CREATE TABLE {target} ({', '.join(columns)}, PRIMARY KEY ({', '.join(keys)}))")
            definitions: dict[str, Any] = {f.name: (get_type_from_str(f.type) | None, None) for f in fields}
            row_model = create_model(f"Imported_{entity}", **definitions)
            frame, first_row = frames[excel]
            seen: set[tuple[Any, ...]] = set()
            for number, (_, raw) in enumerate(frame.iterrows(), first_row):
                mapped = map_row(raw, {f.excel_column: f for f in fields if f.excel_column}, [f.name for f in fields])
                if not any(value not in (None, "") for value in mapped.values()):
                    continue
                try:
                    record = row_model.model_validate(mapped).model_dump()
                except ValidationError as exc:
                    finding(target, "?", f"Typfehler, Zeile ausgelassen: {exc.error_count()} Felder", number)
                    continue
                key = tuple(record[k] for k in keys)
                if not all(key):
                    finding(target, str(key), "Schlüssel fehlt; Zeile ausgelassen", number)
                    continue
                if key in seen:
                    finding(
                        target,
                        str(key),
                        "Schlüssel doppelt; erste Zeile bleibt erhalten",
                        number,
                        {"person": 1, "mandate": 2, "mandate_person": 3}.get(target),
                    )
                    continue
                seen.add(key)
                if record.get("mandate_id") == "SA":
                    finding(
                        target, "SA", "Technischer SA-Auftrag darf nicht im Workbook stehen; Zeile ausgelassen", number
                    )
                    continue
                record.update(source_row=number, is_active=1)
                if target == "person":
                    record["family_key"] = str(record.get("family_id") or "").casefold()
                names = list(record)
                conn.execute(
                    f"INSERT INTO {target} ({','.join(names)}) VALUES ({','.join('?' for _ in names)})",
                    tuple(record.values()),
                )
                total += 1
        conn.execute(
            """INSERT INTO mandate (mandate_id, service_type_id, allowed_travel_time,
                     allowed_direct_effort, allowed_indirect_effort, source_row)
                     VALUES ('SA','ST999',?,?,?,0)""",
            (INTERNAL_ALLOWED_TRAVEL_TIME, INTERNAL_ALLOWED_DIRECT_EFFORT, INTERNAL_ALLOWED_INDIRECT_EFFORT),
        )
        if not conn.execute("SELECT 1 FROM service_types WHERE service_type_id='ST999'").fetchone():
            conn.execute(
                "INSERT INTO service_types (service_type_id,code,hourly_rate,rundung,source_row) VALUES ('ST999','SONST',0,1,0)"
            )
        create_mandate_view(conn)
        for table, refs in FOREIGN_KEYS.items():
            for column, ref_table, ref_column, invariant in refs:
                for key, row, value in conn.execute(
                    f"SELECT t.rowid,t.source_row,t.{column} FROM {table} t LEFT JOIN {ref_table} r ON t.{column}=r.{ref_column} WHERE t.{column} IS NOT NULL AND t.{column}<>'' AND r.{ref_column} IS NULL"
                ):
                    finding(table, str(key), f"{column}={value}: Verweis nicht in {ref_table}", row, invariant)
                    if table == "mandate":
                        conn.execute("UPDATE mandate SET needs_review=1 WHERE rowid=?", (key,))
        for key, row in conn.execute(
            "SELECT mandate_id,source_row FROM v_mandate WHERE mandate_id<>'SA' AND index_person_id IS NULL"
        ):
            finding("mandate", key, "Kein Indexkind bestimmbar; Rechnung kann nicht erstellt werden", row, 8)
        for key, row in conn.execute(
            "SELECT mandate_id,source_row FROM mandate WHERE mandate_id<>'SA' AND COALESCE(application_number,'')=''"
        ):
            finding("mandate", key, "Geschäftsnummer fehlt; Rechnung wird mit PRÜFEN markiert", row, 42)
        for key, row in conn.execute(
            "SELECT mandate_id,source_row FROM mandate WHERE predecessor_mandate_id=mandate_id"
        ):
            finding("mandate", key, "Auftrag ist sein eigener Vorgänger", row, 10)
            conn.execute("UPDATE mandate SET needs_review=1 WHERE mandate_id=?", (key,))
        for (predecessor,) in conn.execute(
            "SELECT predecessor_mandate_id FROM mandate WHERE predecessor_mandate_id IS NOT NULL GROUP BY predecessor_mandate_id HAVING count(*)>1"
        ):
            finding("mandate", predecessor, "Vorgänger hat mehrere Folgeaufträge", invariant=11)
            conn.execute(
                "UPDATE mandate SET needs_review=1 WHERE predecessor_mandate_id=? OR mandate_id=?",
                (predecessor, predecessor),
            )
        for contact, row in conn.execute(
            "SELECT contact_person_id,source_row FROM masterdata_contact_person WHERE COALESCE(notes,'')<>''"
        ):
            finding(
                "masterdata_contact_person",
                contact,
                "Übernommene Ansprechpersonen-Bemerkung prüfen (etwa widersprüchliche Anrede)",
                row,
                13,
            )
            conn.execute("UPDATE mandate SET needs_review=1 WHERE contact_person_id=?", (contact,))
        for key, row, payer, number in conn.execute(
            "SELECT mandate_id,source_row,payer_id,application_number FROM mandate WHERE mandate_id<>'SA'"
        ):
            if payer != "P1000" or not number:
                continue
            if not re.fullmatch(r"[0-9]{9}", number) or number[-3:] == "000":
                finding("mandate", key, "KJA-Geschäftsnummer entspricht nicht yymmddnnn", row, 43)
                conn.execute("UPDATE mandate SET needs_review=1 WHERE mandate_id=?", (key,))
            else:
                try:
                    datetime.strptime(number[:6], "%y%m%d")
                except ValueError:
                    finding("mandate", key, "Ungültiges Datum in KJA-Geschäftsnummer", row, 44)
                    conn.execute("UPDATE mandate SET needs_review=1 WHERE mandate_id=?", (key,))
                if (
                    conn.execute(
                        "SELECT count(*) FROM mandate WHERE payer_id='P1000' AND application_number=?", (number,)
                    ).fetchone()[0]
                    > 1
                ):
                    finding("mandate", key, "KJA-Geschäftsnummer mehrfach verwendet", row, 48)
                    conn.execute("UPDATE mandate SET needs_review=1 WHERE mandate_id=?", (key,))
        for key, row in conn.execute(
            "SELECT person_id,source_row FROM person WHERE family_key<>'' AND date_of_birth IS NULL"
        ):
            finding("person", key, "Kind einer Familie ohne Geburtsdatum", row, 45)
            conn.execute(
                "UPDATE mandate SET needs_review=1 WHERE mandate_id IN (SELECT mandate_id FROM mandate_person WHERE person_id IN (SELECT p.person_id FROM person p JOIN person affected ON p.family_key=affected.family_key WHERE affected.person_id=?))",
                (key,),
            )
        for key, row in conn.execute(
            "SELECT mandate_id,source_row FROM mandate WHERE mandate_id<>'SA' AND contact_person_id IS NULL"
        ):
            finding("mandate", key, "Ansprechperson fehlt", row, 23)
        conn.execute(
            "UPDATE mandate SET needs_review=1 WHERE mandate_id<>'SA' AND (COALESCE(application_number,'')='' OR mandate_id IN (SELECT mandate_id FROM v_mandate WHERE index_person_id IS NULL))"
        )
        for table in ("mandate", "report"):
            for key, row in conn.execute(f"SELECT rowid,source_row FROM {table} WHERE notes LIKE 'ZU PRÜFEN:%'"):
                finding(table, str(key), "Übernommene Angabe zu prüfen", row, 50)
                if table == "mandate":
                    conn.execute("UPDATE mandate SET needs_review=1 WHERE rowid=?", (key,))
    out = config.get_output_path()
    out.mkdir(parents=True, exist_ok=True)
    report = out / "stammdaten_befunde.json"
    report.write_text("[" + ",\n".join(f.model_dump_json() for f in findings) + "]", encoding="utf-8")
    logger.info("Stammdaten importiert: {} Zeilen, {} Befunde ({})", total + 1, len(findings), report)
    if source.parent.resolve() == config.get_imports_path().resolve():
        done = config.get_done_path()
        done.mkdir(parents=True, exist_ok=True)
        shutil.move(str(source), str(done / source.name))
    return total + 1

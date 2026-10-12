"""Generalprüfung vor einem Kundenimport, ohne SQLite oder Quelle zu verändern."""

import hashlib
import re
from pathlib import Path

from openpyxl.utils.cell import get_column_letter, range_boundaries
from openpyxl.utils.datetime import from_excel
from pydantic import BaseModel

from data_imports.normalized_masterdata import OPTIONAL_NEW_COLUMNS, OPTIONAL_TABLES, TABLES
from data_imports.workbook_package import NS, WorkbookPackage, cell_map, tag
from shared_modules.accordix import FIELD_VALUE_LISTS, parse_db_date
from shared_modules.config import Config


class WorkbookIssue(BaseModel):
    """Technischer oder fachlicher Befund ohne Personennamen oder Rohwerte."""

    table: str
    row: int | None = None
    key: str = ""
    field: str | None = None
    message: str
    technical: bool = False


class WorkbookReview(BaseModel):
    """Nachweis der vor dem Import gelesenen Quelle."""

    sha256: str
    record_counts: dict[str, int] = {}
    formula_count: int = 0
    issues: list[WorkbookIssue] = []


def review_workbook(config: Config, source: Path) -> WorkbookReview:
    """Prüft Struktur, Formeln, Schlüssel und Pflichtverweise unabhängig von Caches."""
    package = WorkbookPackage(source)
    result = WorkbookReview(sha256=hashlib.sha256(source.read_bytes()).hexdigest())
    sheet_for_table: dict[str, str] = {}
    roots = {title: package.read_sheet(title) for title in package.sheets}
    for title, path in package.sheets.items():
        shared_anchors = {
            formula.get("si")
            for cell in roots[title].findall(".//m:c", NS)
            if (formula := cell.find(tag("f"))) is not None and formula.get("t") == "shared" and formula.text
        }
        rel_path = path.replace("/worksheets/", "/worksheets/_rels/") + ".rels"
        if rel_path in package.parts:
            from xml.etree import ElementTree as ET

            for relationship in ET.fromstring(package.parts[rel_path]):
                target = relationship.get("Target", "")
                if "/tables/" in target:
                    table_path = "xl/tables/" + target.rsplit("/", 1)[-1]
                    sheet_for_table[ET.fromstring(package.parts[table_path]).get("displayName", "")] = title
        for cell in roots[title].findall(".//m:c", NS):
            if cell.find(tag("f")) is not None:
                formula_node = cell.find(tag("f"))
                assert formula_node is not None
                if formula_node.get("t") == "shared" and formula_node.get("si") not in shared_anchors:
                    result.issues.append(
                        WorkbookIssue(
                            table=title,
                            field=cell.get("r"),
                            message="Excel-Sammelformel ohne Basisformel",
                            technical=True,
                        )
                    )
                result.formula_count += 1
                formula = cell.findtext(tag("f"), "")
                if "#REF!" in formula or cell.get("t") == "e":
                    result.issues.append(
                        WorkbookIssue(
                            table=title, field=cell.get("r"), message="Formelfehler oder defekter Bezug", technical=True
                        )
                    )
                if cell.find(tag("v")) is None:
                    result.issues.append(
                        WorkbookIssue(
                            table=title,
                            field=cell.get("r"),
                            message="Formelwert fehlt; vor Import neu berechnen",
                            technical=True,
                        )
                    )
    records: dict[str, list[tuple[int, dict[str, str]]]] = {}
    table_map = TABLES | {name: values for name, values in OPTIONAL_TABLES.items() if values[1] in config.models}
    for name, (_, entity) in table_map.items():
        if name not in sheet_for_table:
            if name not in OPTIONAL_TABLES:
                result.issues.append(WorkbookIssue(table=name, message="Pflichttabelle fehlt", technical=True))
            continue
        _, table = package.table(name)
        headers = [column.get("name", "") for column in table.findall("m:tableColumns/m:tableColumn", NS)]
        fields = config.models[entity].fields
        for field in fields:
            if field.excel_column not in headers and not (
                field.optional and field.excel_column in OPTIONAL_NEW_COLUMNS
            ):
                result.issues.append(
                    WorkbookIssue(table=name, field=field.excel_column, message="Importspalte fehlt", technical=True)
                )
        left, top, _, bottom = range_boundaries(table.get("ref", ""))
        assert left and top and bottom
        cells = cell_map(roots[sheet_for_table[name]])
        seen: set[tuple[str, ...]] = set()
        actual: list[tuple[int, dict[str, str]]] = []
        for row in range(top + 1, bottom + 1):
            values = {
                field.name: package.value(
                    cells.get(f"{get_column_letter(left + headers.index(field.excel_column))}{row}")
                )
                for field in fields
                if field.excel_column in headers
            }
            if not any(values.values()):
                continue
            actual.append((row, values))
            key = tuple(values.get(field.name, "") for field in fields if field.primary_key)
            duplicate_key = key
            if name == "relation_mandate_emp":
                duplicate_key += tuple(values.get(field, "") for field in ("role", "valid_from", "valid_to"))
            if not all(key):
                result.issues.append(WorkbookIssue(table=name, row=row, message="Schlüssel fehlt"))
            elif duplicate_key in seen:
                result.issues.append(
                    WorkbookIssue(
                        table=name,
                        row=row,
                        key="/".join(key),
                        message="Schlüssel doppelt; Import würde weitere Zeile auslassen",
                    )
                )
            seen.add(duplicate_key)
            for field in fields:
                if not field.optional and not values.get(field.name):
                    result.issues.append(
                        WorkbookIssue(
                            table=name, row=row, key="/".join(key), field=field.name, message="Pflichtangabe fehlt"
                        )
                    )
            dates = {}
            for field in (
                "date_of_birth",
                "due_date",
                "completed_date",
                "start_date",
                "end_date",
                "valid_from",
                "valid_to",
            ):
                value = values.get(field)
                if not value:
                    continue
                parsed = parse_db_date(value)
                if parsed is None:
                    try:
                        converted = from_excel(float(value))
                        from datetime import datetime

                        parsed = converted.date() if isinstance(converted, datetime) else None
                    except (ValueError, OverflowError):
                        parsed = None
                dates[field] = parsed
                if parsed is None:
                    result.issues.append(
                        WorkbookIssue(
                            table=name, row=row, key="/".join(key), field=field, message="Datumsangabe ist ungültig"
                        )
                    )
            first = dates.get("start_date", dates.get("valid_from"))
            last = dates.get("end_date", dates.get("valid_to"))
            if first and last and first > last:
                result.issues.append(
                    WorkbookIssue(table=name, row=row, key="/".join(key), message="Enddatum liegt vor Beginn")
                )
            for field, allowed in FIELD_VALUE_LISTS.items():
                if field == "gender" and name != "person":
                    continue
                if values.get(field) and values[field] not in allowed:
                    result.issues.append(
                        WorkbookIssue(
                            table=name,
                            row=row,
                            key="/".join(key),
                            field=field,
                            message="Codewert weicht von Werteliste ab",
                        )
                    )
        records[name] = actual
        result.record_counts[name] = len(actual)
        input_columns = {
            get_column_letter(left + headers.index(field.excel_column))
            for field in fields
            if field.excel_column in headers
        }
        for reference, cell in cells.items():
            column, number = re.findall(r"^([A-Z]+)(\d+)$", reference)[0]
            if int(number) > bottom and column in input_columns and cell.find(tag("f")) is None and package.value(cell):
                result.issues.append(
                    WorkbookIssue(
                        table=name,
                        row=int(number),
                        field=column,
                        message="Eingabe liegt ausserhalb der Excel-Tabelle und würde nicht importiert",
                        technical=True,
                    )
                )
    from data_imports.normalized_masterdata import FOREIGN_KEYS

    physical = {values[0]: name for name, values in table_map.items()}
    for table, links in FOREIGN_KEYS.items():
        excel = physical.get(table, table)
        for row, record in records.get(excel, []):
            for field, target, key, _ in links:
                value = record.get(field)
                target_records = records.get(physical.get(target, target), [])
                if value and value not in {entry.get(key) for _, entry in target_records}:
                    result.issues.append(
                        WorkbookIssue(table=excel, row=row, field=field, message="Verweis ist unbekannt")
                    )
    return result

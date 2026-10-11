"""Neu erzeugte Mitarbeiterblätter mit zweijährigem Berichtsraster."""

import re
import sqlite3
from collections import defaultdict
from datetime import date
from pathlib import Path

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from reports.report_data import Employee, ReportAppointment, assignment_problem, load_report_data
from shared_modules.config import Config
from shared_modules.month_period import get_month_period
from shared_modules.utils import ensure_dir

_MONTHS = ("Jan", "Feb", "Mrz", "Apr", "Mai", "Jun", "Jul", "Aug", "Sept", "Okt", "Nov", "Dez")
_COLOURS = {"Bericht": "FFFF0000", "Zwischenbericht": "FFFFC000", "Abschlussbericht": "FF00B0F0"}
_HEADERS = (
    "Auftrag",
    "Vorname",
    "Nachname",
    "Leistungsbesteller",
    "Ansprechperson",
    "Telefon",
    "Beginn",
    "Ende",
    "Zuordnung aktuell",
    "Fahrt h/Monat",
    "Direkt h/Monat",
    "Indirekt h/Monat",
    "Gesamt h/Monat",
    "Leistung",
)


def create_klientenuebersicht(config: Config, reporting_month: str) -> Path:
    """Erzeugt einen Auftragsstand ohne Rückschreiben und ohne Sammelblatt Übersicht."""
    period = get_month_period(reporting_month)
    year, month = period.start.year, period.start.strftime("%Y-%m")
    if year >= 9999:
        raise ValueError("Für die Zweijahresübersicht muss das Folgejahr darstellbar sein.")
    db = config.get_db_path()
    if not db.is_file():
        raise FileNotFoundError(f"SQLite-Datenbank fehlt: {db}. Zuerst import-master ausführen.")
    with sqlite3.connect(db) as conn:
        employees, reports = load_report_data(conn)
    employee_ids = {employee.emp_id for employee in employees}
    grouped: dict[str, list[ReportAppointment]] = {employee.emp_id: [] for employee in employees}
    findings: list[tuple[ReportAppointment, str]] = []
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    _legend(workbook, month, year)
    for report in reports:
        problem = assignment_problem(report, employee_ids)
        if problem:
            report.assignment_error = problem
            findings.append((report, problem))
            logger.warning(
                "Klientenübersicht: Bericht {} / Auftrag {}: {}", report.report_id, report.mandate_id, problem
            )
        elif report.due_date and year <= report.due_date.year <= year + 1:
            assert report.employee_id is not None
            grouped[report.employee_id].append(report)
        else:
            findings.append((report, "Fälligkeit ausserhalb des Zweijahresrasters"))
    for employee in employees:
        _employee_sheet(workbook, employee, grouped[employee.emp_id], month, year)
    _details(workbook, reports)
    _findings(workbook, findings)
    output = ensure_dir(config.get_output_path() / f"Klientenuebersichten_{month}")
    path = output / f"Klientenuebersicht_{month}.xlsx"
    workbook.save(path)
    workbook.close()
    logger.info("Klientenübersicht geschrieben: {} ({} Berichte, {} Befunde)", path, len(reports), len(findings))
    return path


def _cell(sheet: Worksheet, row: int, column: int, value: str | int | float | date | None) -> None:
    """Freitext bleibt Text, auch bei führendem Formelzeichen."""
    cell = sheet.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.font = Font(name="Arial", size=10)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if isinstance(value, date):
        cell.number_format = "dd.mm.yyyy"


def _header(sheet: Worksheet, row: int, values: tuple[str, ...]) -> None:
    """Beschriftet einen filterbaren Tabellenkopf."""
    for column, value in enumerate(values, 1):
        _cell(sheet, row, column, value)
        cell = sheet.cell(row, column)
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4057")
    sheet.row_dimensions[row].height = 34


def _legend(workbook: Workbook, month: str, year: int) -> None:
    """Erklärt ausschliesslich die tatsächlich erzeugten Daten und Farben."""
    sheet = workbook.create_sheet("Legende")
    _header(sheet, 1, ("Klientenübersicht", "Bedeutung"))
    lines = [
        ("Auftragsstand", f"Aus der importierten Datenbank neu erzeugt für {month}; Raster {year}/{year + 1}."),
        (
            "Datenpflege",
            "Diese Datei ist eine Ausgabe. Änderungen in der Stammdatenmappe pflegen, DB neu aufbauen und Report neu erzeugen. Keine zweite manuell gepflegte Datenquelle.",
        ),
        (
            "Zeilen",
            "Eine Zeile je Auftrag mit Berichtsterminen im Raster. Parallel- und Folgeaufträge bleiben getrennt; Auftrag-Nr steht in A. Aufträge ohne Berichtseinträge werden nicht ergänzt.",
        ),
        (
            "Spalten G:N",
            "Aktuell importierte Angaben zum angezeigten Auftrag. Beginn/Ende = Bewilligung. Zuordnung aktuell = heutige Zuordnung MA, unabhängig von der datierten Berichtszuständigkeit.",
        ),
        (
            "Kontingente J:M",
            "Bewilligte Fahrtzeit, direkter und indirekter Fallkontakt in Stunden pro Monat; Gesamt = Summe. Keine geleisteten Stunden.",
        ),
        (
            "Zuständigkeit",
            "Rolle P am Fälligkeitstag aus Zuständigkeitsverlauf. Ohne Historie: aktuelle P-Zuordnung, historische Zuständigkeit unbekannt. Lücken/Mehrfach-P erscheinen als Befund.",
        ),
        ("Bericht", "Rot"),
        ("Zwischenbericht", "Orange"),
        ("Abschlussbericht", "Blau; keine Kombination mit einer zweiten Berichtsform."),
        (
            "Gemischte Berichtsformen",
            "Graue Zelle mit einzeln beschrifteten Terminen. Mehrere Termine werden einzeln aufgeführt.",
        ),
        (
            "Status",
            "Alle eingetragenen Berichte werden angezeigt, auch erledigt/entfällt. Ein leerer Status bedeutet nicht erledigt.",
        ),
        (
            "Berichtsdetails",
            "Vollständige Liste aller expliziten Berichtseinträge, auch ohne Datum oder ausserhalb des Rasters. Keine Termine aus Berichtsintervallen erzeugt.",
        ),
        ("Nicht zugeordnet", "Befunde und Termine ausserhalb des Rasters. Diese Liste vor Verwendung prüfen."),
    ]
    for row, (label, explanation) in enumerate(lines, 2):
        _cell(sheet, row, 1, label)
        _cell(sheet, row, 2, explanation)
        if label in _COLOURS:
            sheet.cell(row, 1).fill = PatternFill("solid", fgColor=_COLOURS[label])
        sheet.row_dimensions[row].height = 45
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 105
    sheet.freeze_panes = "B2"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    sheet.print_area = f"A1:B{sheet.max_row}"


def _employee_sheet(
    workbook: Workbook, employee: Employee, reports: list[ReportAppointment], month: str, year: int
) -> None:
    """Erhält Stammdaten links und zwei Monatsblöcke ohne wiederholte Namensspalten."""
    title = re.sub(r"[\\/*?:\[\]]", "_", f"{employee.emp_id} {employee.first_name or employee.last_name or 'MA'}")[:31]
    base = title.strip(" '") or "MA"
    title = base
    suffix = 1
    while title.casefold() in {name.casefold() for name in workbook.sheetnames}:
        suffix += 1
        title = f"{base[:26]}_{suffix}"
    sheet = workbook.create_sheet(title)
    _cell(sheet, 2, 1, f"{employee.first_name or ''} {employee.last_name or ''} ({employee.emp_id})".strip())
    sheet.merge_cells("A2:N2")
    for column in range(15, 39, 3):
        _cell(sheet, 2, column, str(year if column < 27 else year + 1))
        sheet.merge_cells(start_row=2, end_row=2, start_column=column, end_column=column + 2)
    _cell(
        sheet,
        3,
        1,
        f"Auftragsstand {month}; neu erzeugte Ausgabe. Berichtszuständigkeit nach Fälligkeit; siehe Legende.",
    )
    sheet.merge_cells("A3:N3")
    sheet.row_dimensions[2].height = 28
    sheet.row_dimensions[3].height = 24
    _header(sheet, 4, _HEADERS + _MONTHS + _MONTHS)
    for column in range(15, 39):
        _cell(sheet, 5, column, str(year if column < 27 else year + 1))
    by_mandate: dict[str, list[ReportAppointment]] = defaultdict(list)
    for report in reports:
        assert report.mandate_id is not None
        by_mandate[report.mandate_id].append(report)
    for row, (_, appointments) in enumerate(
        sorted(by_mandate.items(), key=lambda item: (item[1][0].last_name or "", item[1][0].first_name or "", item[0])),
        6,
    ):
        report = appointments[0]
        minutes = (report.travel_minutes, report.direct_minutes, report.indirect_minutes)
        values = (
            report.mandate_id,
            report.first_name,
            report.last_name,
            report.requester,
            report.contact,
            report.phone,
            report.start_date,
            report.end_date,
            report.current_employees,
            *(value / 60 if value is not None else None for value in minutes),
            sum(value or 0 for value in minutes) / 60 if any(value is not None for value in minutes) else None,
            report.service_type,
        )
        for column, value in enumerate(values, 1):
            _cell(sheet, row, column, value)
            if 10 <= column <= 13:
                sheet.cell(row, column).number_format = "0.00"
        by_month: dict[int, list[ReportAppointment]] = defaultdict(list)
        for appointment in appointments:
            assert appointment.due_date is not None
            column = 15 + (appointment.due_date.year - year) * 12 + appointment.due_date.month - 1
            by_month[column].append(appointment)
        height = 40.0
        for column, entries in by_month.items():
            entries.sort(key=lambda item: (item.due_date or date.min, item.report_id))
            text = "\n".join(
                f"{entry.due_date:%d.%m.} {entry.report_form or 'Form fehlt'} [{entry.report_id}] ({entry.status or 'Status fehlt'})"
                for entry in entries
            )
            _cell(sheet, row, column, text)
            forms = {entry.report_form for entry in entries}
            colour = _COLOURS.get(next(iter(forms)) or "", "FFE2E8F0") if len(forms) == 1 else "FFE2E8F0"
            sheet.cell(row, column).fill = PatternFill("solid", fgColor=colour)
            height = max(height, 44.0 * len(entries))
        sheet.row_dimensions[row].height = min(height, 409)
    if not reports:
        _cell(sheet, 6, 1, "Keine Berichtstermine im Raster.")
        sheet.merge_cells("A6:N6")
    widths = (14, 16, 18, 26, 25, 20, 12, 12, 20, 11, 11, 11, 11, 24) + (25,) * 24
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.freeze_panes = "O6"
    sheet.auto_filter.ref = f"A4:AL{max(sheet.max_row, 6)}"
    sheet.print_title_rows = "2:5"
    sheet.print_title_cols = "A:N"
    sheet.print_area = f"A2:{'AL' if reports else 'N'}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    # Drei Monate je Druckseite, mit Stammdaten links auf jeder Seite.
    sheet.page_setup.scale = 55
    if reports:
        for column in range(17, 38, 3):
            sheet.col_breaks.append(Break(id=column))


def _details(workbook: Workbook, reports: list[ReportAppointment]) -> None:
    """Bewahrt jeden Bericht genau einmal, unabhängig vom Raster und dessen Status."""
    sheet = workbook.create_sheet("Berichtsdetails")
    _header(
        sheet,
        1,
        (
            "Bericht-ID",
            "Auftrag",
            "Fälligkeit",
            "Berichtsform",
            "Status",
            "Erledigt am",
            "Zuständig",
            "Zuordnungsbasis",
            "Befund",
            "Notizen",
        ),
    )
    for row, report in enumerate(reports, 2):
        for column, value in enumerate(
            (
                report.report_id,
                report.mandate_id,
                report.due_date,
                report.report_form,
                report.status,
                report.completed_date,
                report.employee_id,
                report.assignment_basis,
                report.assignment_error,
                report.notes,
            ),
            1,
        ):
            _cell(sheet, row, column, value)
    _list_layout(sheet, (14, 14, 14, 22, 16, 14, 16, 48, 55, 55))


def _findings(workbook: Workbook, findings: list[tuple[ReportAppointment, str]]) -> None:
    """Führt unzuordenbare und ausserhalb liegende Termine sichtbar auf."""
    sheet = workbook.create_sheet("Nicht zugeordnet")
    _header(sheet, 1, ("Bericht-ID", "Auftrag", "Fälligkeit", "Befund"))
    for row, (report, reason) in enumerate(findings, 2):
        for column, value in enumerate((report.report_id, report.mandate_id, report.due_date, reason), 1):
            _cell(sheet, row, column, value)
    if not findings:
        _cell(sheet, 2, 1, "Keine Befunde.")
    _list_layout(sheet, (14, 14, 14, 80))


def _list_layout(sheet: Worksheet, widths: tuple[int, ...]) -> None:
    """Formatiert die Kontrolllisten für Filter und Ausdruck."""
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 48
    sheet.freeze_panes = "C2"
    end = get_column_letter(len(widths))
    sheet.auto_filter.ref = f"A1:{end}{sheet.max_row}"
    sheet.print_title_rows = "1:1"
    sheet.print_area = f"A1:{end}{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)

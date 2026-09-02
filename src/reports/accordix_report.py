"""
Erstellt die KFSG-Leistungsmeldung für die Datenbank Accordix (Kanton Bern, KJA).

Die Meldung wird auf Basis der offiziellen Excel-Vorlage
«Import-Accordix_ambulant_Excel-Format_V1.0_DE.xlsx» erzeugt, damit das
vorgegebene Format strikt eingehalten wird (Kopfzeilen 1-6, Werte-Blatt,
Wertelisten). Datenzeilen beginnen ab Zeile 7.

Quelle der Formatvorgaben:
https://www.kja.dij.be.ch/de/start/foerder--und-schutzleistungen/kantonale-datenerfassung/Datenbank_Accordix/MoeglichkeitenDatenmeldung.html

Befüllt werden alle Felder, die in den Klienten-Stammdaten gepflegt sind:
Nachname, Vorname, AHV-Nummer, Geburtsdatum, Geschlecht, UMA/UMF,
Hauptsprache, Wohnkanton, Wohnort Sorgeberechtigte, Leistungsart,
Zuweisung, Eintrittsdatum sowie das Austrittsdatum (nur wenn das
Leistungsende im oder vor dem Meldemonat liegt).
Die Austrittsfelder (war geplant, Austrittsgrund, Situation nach Austritt)
sind Ereignisdaten und werden im Meldefile manuell nachgetragen.
"""

import sqlite3
from datetime import date
from pathlib import Path
from typing import Any, Optional

from loguru import logger
from openpyxl import load_workbook

from shared_modules.accordix import (
    FIELD_VALUE_LISTS,
    SERVICE_TYPE_MAP,
    missing_required_fields,
    parse_db_date,
    validate_coded_value,
)
from shared_modules.config import Config
from shared_modules.internal_client import INTERNAL_SERVICE_TYPE_ID
from shared_modules.month_period import get_month_period
from shared_modules.utils import ensure_dir

_TEMPLATE_NAME = "Import-Accordix_ambulant_Excel-Format_V1.0_DE.xlsx"
_SHEET_NAME = "Ambulant"
_FIRST_DATA_ROW = 7  # Kopfzeilen/Hinweise stehen in den Zeilen 1-6

# Spalten gemäss Vorlage (Zeile 4 enthält die technischen Feldnamen)
_COL_LAST_NAME = "A"  # LastName
_COL_FIRST_NAME = "B"  # FirstName
_COL_AHV = "C"  # SocialInsuranceNumber
_COL_DATE_OF_BIRTH = "D"  # DateOfBirth
_COL_GENDER = "E"  # Gender
_COL_UMA_UMF = "F"  # UmaUmfRecognition
_COL_LANGUAGE = "G"  # SpokenLanguage
_COL_CANTON = "I"  # CantonOfResidence
_COL_RESIDENCE = "J"  # ResidenceLegalGuardian
_COL_SERVICE_TYPE = "L"  # ServiceTypeName
_COL_ALLOCATION = "M"  # Allocation
_COL_START_DATE = "Q"  # StartDate
_COL_END_DATE = "S"  # EndDate

# Neue Stammdatenfelder, die beim letzten Import ergänzt wurden
_ACCORDIX_DB_COLUMNS = (
    "date_of_birth",
    "gender",
    "uma_umf",
    "spoken_language",
    "canton_of_residence",
    "residence_legal_guardian",
    "allocation",
)

_SQL_BASE = """
SELECT
    c.client_id,
    c.last_name,
    c.first_name,
    c.social_security_number,
    c.start_date,
    c.end_date,
    st.code AS service_code
    {extra_cols}
FROM clients c
LEFT JOIN service_types st ON c.service_type = st.service_type_id
WHERE date(c.start_date) <= date(:period_end)
  AND (c.end_date IS NULL OR date(c.end_date) >= date(:period_start))
  AND COALESCE(c.service_type, '') <> :internal_service_type
ORDER BY c.last_name, c.first_name
"""


def _format_date(value: date) -> str:
    """Formatiert ein Datum gemäss Vorlage als TT.MM.JJJJ (Text)."""
    return value.strftime("%d.%m.%Y")


def _existing_accordix_columns(conn: sqlite3.Connection) -> list[str]:
    """Liefert die Accordix-Spalten, die in der Tabelle clients bereits existieren."""
    existing = {row[1] for row in conn.execute("PRAGMA table_info(clients)").fetchall()}
    return [col for col in _ACCORDIX_DB_COLUMNS if col in existing]


def _set_if_present(sheet, column: str, excel_row: int, value: Optional[Any]) -> None:
    """Schreibt einen Wert in eine Zelle, sofern er nicht leer ist."""
    if value is not None and str(value).strip() != "":
        sheet[f"{column}{excel_row}"] = str(value).strip()


def create_accordix_report(config: Config, reporting_month: str) -> Path:
    """
    Erstellt die Accordix-Leistungsmeldung (ambulant) für einen Meldemonat.

    Enthalten sind alle Klient:innen, deren Leistung im Meldemonat aktiv war
    (start_date <= Monatsende und end_date offen oder >= Monatsanfang), deren
    Leistungsart auf eine Accordix-Leistungsart abbildbar ist und deren
    Accordix-Pflichtfelder vollständig sind. Unvollständige Datensätze werden
    mit Angabe der fehlenden Felder übersprungen.

    Args:
        config: Konfigurationsobjekt.
        reporting_month: Meldemonat (beliebiges Format MM.YYYY / MM-YYYY / YYYY-MM).

    Returns:
        Pfad zur erzeugten Excel-Datei.
    """
    period = get_month_period(reporting_month)
    month_str = period.start.strftime("%Y-%m")
    period_start = period.start.date()
    period_end = period.end.date()

    template_path = config.get_template_path(_TEMPLATE_NAME)
    if not template_path.exists():
        raise FileNotFoundError(f"Accordix-Vorlage nicht gefunden: {template_path}")

    db_path = config.get_db_path()
    with sqlite3.connect(db_path) as conn:
        accordix_cols = _existing_accordix_columns(conn)
        missing_cols = [col for col in _ACCORDIX_DB_COLUMNS if col not in accordix_cols]
        if missing_cols:
            logger.warning(
                "Spalten fehlen in der Tabelle clients: {}. "
                "Bitte zuerst 'import-master' mit der erweiterten Stammdaten-Datei ausführen.",
                ", ".join(missing_cols),
            )
        extra_cols = "".join(f",\n    c.{col}" for col in accordix_cols)
        cursor = conn.execute(
            _SQL_BASE.format(extra_cols=extra_cols),
            {
                "period_start": period_start.isoformat(),
                "period_end": period_end.isoformat(),
                "internal_service_type": INTERNAL_SERVICE_TYPE_ID,
            },
        )
        column_names = [desc[0] for desc in cursor.description]
        rows = [dict(zip(column_names, row)) for row in cursor.fetchall()]

    if not rows:
        logger.warning("Keine aktiven Klient:innen für {} gefunden.", month_str)
        raise ValueError(f"Keine aktiven Klient:innen für {month_str} gefunden.")

    # load_workbook verwirft die x14-Datenvalidierungs-Erweiterung der Vorlage
    # (Dropdown-Komfort); das strikt einzuhaltende Importformat (Zellinhalte)
    # bleibt davon unberührt.
    workbook = load_workbook(template_path)
    sheet = workbook[_SHEET_NAME]

    written = 0
    skipped = 0
    for row in rows:
        client_id = row["client_id"]
        label = f"{row['last_name']} {row['first_name']}"

        service_name = SERVICE_TYPE_MAP.get(row["service_code"] or "")
        if service_name is None:
            skipped += 1
            logger.warning(
                "Klient {} ({}) übersprungen: Leistungsart {!r} ist keiner "
                "Accordix-Leistungsart zugeordnet.",
                client_id,
                label,
                row["service_code"],
            )
            continue

        start_date = parse_db_date(row["start_date"])
        if start_date is None:
            skipped += 1
            logger.warning(
                "Klient {} ({}) übersprungen: kein gültiges Eintrittsdatum.",
                client_id,
                label,
            )
            continue

        missing = missing_required_fields(row)
        if missing:
            skipped += 1
            logger.warning(
                "Klient {} ({}) übersprungen: fehlende Accordix-Pflichtfelder: {}. "
                "Bitte in den Stammdaten nachtragen und 'import-master' ausführen.",
                client_id,
                label,
                ", ".join(missing),
            )
            continue

        invalid = [
            f"{field}: {error}"
            for field in FIELD_VALUE_LISTS
            if (error := validate_coded_value(field, row.get(field))) is not None
        ]
        if invalid:
            skipped += 1
            logger.warning(
                "Klient {} ({}) übersprungen: ungültige Werte: {}. "
                "Bitte in den Stammdaten korrigieren und 'import-master' ausführen.",
                client_id,
                label,
                "; ".join(invalid),
            )
            continue

        end_date = parse_db_date(row["end_date"])

        excel_row = _FIRST_DATA_ROW + written
        sheet[f"{_COL_LAST_NAME}{excel_row}"] = row["last_name"]
        sheet[f"{_COL_FIRST_NAME}{excel_row}"] = row["first_name"]
        _set_if_present(sheet, _COL_AHV, excel_row, row["social_security_number"])
        date_of_birth = parse_db_date(row.get("date_of_birth"))
        if date_of_birth is not None:
            sheet[f"{_COL_DATE_OF_BIRTH}{excel_row}"] = _format_date(date_of_birth)
        _set_if_present(sheet, _COL_GENDER, excel_row, row.get("gender"))
        _set_if_present(sheet, _COL_UMA_UMF, excel_row, row.get("uma_umf"))
        _set_if_present(sheet, _COL_LANGUAGE, excel_row, row.get("spoken_language"))
        _set_if_present(sheet, _COL_CANTON, excel_row, row.get("canton_of_residence"))
        _set_if_present(sheet, _COL_RESIDENCE, excel_row, row.get("residence_legal_guardian"))
        sheet[f"{_COL_SERVICE_TYPE}{excel_row}"] = service_name
        _set_if_present(sheet, _COL_ALLOCATION, excel_row, row.get("allocation"))
        sheet[f"{_COL_START_DATE}{excel_row}"] = _format_date(start_date)
        # Austrittsdatum nur setzen, wenn die Leistung im oder vor dem
        # Meldemonat endet; andernfalls läuft die Leistung weiter (leer).
        if end_date is not None and end_date <= period_end:
            sheet[f"{_COL_END_DATE}{excel_row}"] = _format_date(end_date)
            logger.warning(
                "Klient {} ({}): Austrittsdatum {} gesetzt — die Austrittsfelder "
                "(war geplant, Austrittsgrund, Situation nach Austritt) müssen "
                "im Meldefile manuell ergänzt werden.",
                client_id,
                label,
                _format_date(end_date),
            )
        written += 1

    if written == 0:
        raise ValueError(
            f"Keine meldepflichtigen Leistungen für {month_str} gefunden "
            f"({skipped} Klient:innen übersprungen)."
        )

    output_path = ensure_dir(config.get_output_path())
    out_file = output_path / f"Accordix_ambulant_{month_str}.xlsx"
    workbook.save(out_file)

    logger.info(
        "Accordix-Meldung geschrieben: {} ({} Zeilen, {} übersprungen)",
        out_file,
        written,
        skipped,
    )
    return out_file

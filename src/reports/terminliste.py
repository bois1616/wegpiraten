"""Monatliche Terminlisten für die zuständigen Mitarbeitenden."""

import sqlite3
from pathlib import Path
from urllib.parse import quote

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties

from reports.report_data import Employee, ReportAppointment, assignment_problem, load_report_data
from shared_modules.config import Config
from shared_modules.month_period import get_month_period
from shared_modules.utils import ensure_dir


def create_terminlisten(config: Config, reporting_month: str) -> list[Path]:
    """Erstellt je Mitarbeitenden eine Excel-Datei; unzuordenbare Termine nur melden."""
    period = get_month_period(reporting_month)
    month = period.start.strftime("%Y-%m")
    db_path = config.get_db_path()
    if not db_path.is_file():
        raise FileNotFoundError(f"SQLite-Datenbank fehlt: {db_path}. Zuerst import-master ausführen.")
    with sqlite3.connect(db_path) as conn:
        employees, reports = load_report_data(conn)
        appointments = [
            report
            for report in reports
            if report.due_date is not None
            and period.start.date() <= report.due_date <= period.end.date()
            and report.status not in ("erledigt", "entfällt")
        ]
    grouped: dict[str, list[ReportAppointment]] = {employee.emp_id: [] for employee in employees}
    skipped = 0
    for appointment in appointments:
        reason = assignment_problem(appointment, set(grouped))
        if reason:
            logger.warning(
                "Terminliste: Bericht {} / Auftrag {} ausgelassen: {}",
                appointment.report_id,
                appointment.mandate_id,
                reason,
            )
            skipped += 1
            continue
        assert appointment.employee_id is not None
        grouped[appointment.employee_id].append(appointment)

    output = ensure_dir(config.get_output_path() / f"Terminlisten_{month}")
    files: list[Path] = []
    for employee in employees:
        path = output / f"Terminliste_{month}_{quote(employee.emp_id, safe='')}.xlsx"
        _write_excel(employee, grouped[employee.emp_id], path, month)
        files.append(path)
        logger.info("Terminliste geschrieben: {} ({} Termine)", path, len(grouped[employee.emp_id]))
    logger.info(
        "Terminlisten {}: {} Dateien, {} Termine, {} ausgelassen",
        month,
        len(files),
        len(appointments) - skipped,
        skipped,
    )
    return files


def _write_excel(employee: Employee, appointments: list[ReportAppointment], path: Path, month: str) -> None:
    """Schreibt eine filterbare und druckbare Liste mit echten Excel-Datumswerten."""
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Termine"
    sheet.append([f"Terminliste {month}"])
    sheet.append([f"{employee.last_name or ''} {employee.first_name or ''} ({employee.emp_id})".strip()])
    sheet.append(["Offene Berichte und Einträge ohne Status, fällig in diesem Monat."])
    sheet.append(
        [
            "Fällig am",
            "Auftrag",
            "Indexkind",
            "Leistungsart",
            "Leistungsbesteller",
            "Berichtsform",
            "Status",
            "Notizen",
            "Bericht-ID",
            "Zuordnungsbasis",
        ]
    )
    for appointment in appointments:
        sheet.append(
            [
                appointment.due_date,
                appointment.mandate_id,
                appointment.index_child,
                appointment.service_type,
                appointment.requester,
                appointment.report_form,
                appointment.status,
                appointment.notes,
                appointment.report_id,
                appointment.assignment_basis,
            ]
        )
    if not appointments:
        sheet.append(["Keine fälligen Berichte."])
    # Freitext bleibt Text, auch wenn er mit einem Formelzeichen beginnt.
    for row in sheet:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet["A1"].font = Font(name="Arial", size=16, bold=True)
    for cell in sheet[4]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2E4057")
    for row_number in range(5, 5 + len(appointments)):
        sheet.cell(row_number, 1).number_format = "dd.mm.yyyy"
    for column, width in enumerate((14, 14, 26, 20, 26, 22, 14, 45, 14, 40), 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    for row_number in (1, 2, 3):
        sheet.merge_cells(start_row=row_number, start_column=1, end_row=row_number, end_column=10)
    sheet.freeze_panes = "C5"
    sheet.auto_filter.ref = f"A4:J{max(4, 4 + len(appointments))}"
    sheet.print_title_rows = "1:4"
    sheet.print_area = f"A1:J{sheet.max_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    workbook.save(path)
    workbook.close()

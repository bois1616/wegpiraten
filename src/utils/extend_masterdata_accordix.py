"""
Erweitert die Stammdaten-Datei (wegpiraten_datenbank.xlsx) um die Accordix-Felder.

Die Erweiterung ist idempotent und umfasst:
- neues Blatt «Wertelisten» mit den Accordix-Wertelisten (benannte Bereiche)
- sieben neue Spalten in der Excel-Tabelle «masterdata_client»
  (date_of_birth, gender, uma_umf, spoken_language, canton_of_residence,
  residence_legal_guardian, allocation)
- Dropdown-Datenvalidierungen auf den Code-Spalten (verknüpft mit den
  benannten Bereichen in «Wertelisten», vermeidet Tippfehler und
  Mehrfachpflege der Listen)
- Datumsformat TT.MM.JJJJ für die Spalte date_of_birth

Hinweis: Die erweiterte Datei muss anschliessend wieder nach Proton Drive
hochgeladen werden, damit die Struktur dauerhaft erhalten bleibt.
"""

from pathlib import Path

from loguru import logger
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import TableColumn

from shared_modules.accordix import (
    ACCORDIX_CLIENT_FIELDS,
    FIELD_VALUE_LISTS,
)

CLIENT_TABLE_NAME = "masterdata_client"
VALUELIST_SHEET_NAME = "Wertelisten"
DATE_COLUMN = "date_of_birth"
DATE_FORMAT = "DD.MM.YYYY"
# Datenzeilen, für die Dropdowns und Formate vorbelegt werden
_FIRST_DATA_ROW = 3
_LAST_VALIDATED_ROW = 500

_PREFIX = "accordix_"


def _ensure_valuelist_sheet(workbook) -> None:
    """Legt das Blatt «Wertelisten» mit benannten Bereichen an bzw. aktualisiert es."""
    if VALUELIST_SHEET_NAME in workbook.sheetnames:
        sheet = workbook[VALUELIST_SHEET_NAME]
        logger.info("Blatt '{}' existiert bereits, Wertelisten werden geprüft.", VALUELIST_SHEET_NAME)
    else:
        sheet = workbook.create_sheet(VALUELIST_SHEET_NAME)
        logger.info("Blatt '{}' angelegt.", VALUELIST_SHEET_NAME)

    for col_idx, (field, values) in enumerate(FIELD_VALUE_LISTS.items(), start=1):
        col_letter = get_column_letter(col_idx)
        sheet[f"{col_letter}1"] = field
        for row_idx, value in enumerate(values, start=2):
            sheet[f"{col_letter}{row_idx}"] = value
        sheet.column_dimensions[col_letter].width = max(len(field), max(len(v) for v in values)) + 2

        # Benannten Bereich setzen (bestehenden ersetzen)
        name = f"{_PREFIX}{field}"
        ref = f"{VALUELIST_SHEET_NAME}!${col_letter}$2:${col_letter}${len(values) + 1}"
        if name in workbook.defined_names:
            del workbook.defined_names[name]
        workbook.defined_names[name] = DefinedName(name, attr_text=ref)


def _find_client_table(workbook):
    """Liefert (Sheet, Tabelle) der Excel-Tabelle masterdata_client."""
    for ws in workbook.worksheets:
        if CLIENT_TABLE_NAME in ws.tables:
            return ws, ws.tables[CLIENT_TABLE_NAME]
    raise ValueError(f"Excel-Tabelle '{CLIENT_TABLE_NAME}' nicht gefunden.")


def _ensure_client_columns(workbook) -> None:
    """Erweitert die Tabelle masterdata_client um die Accordix-Spalten."""
    sheet, table = _find_client_table(workbook)
    header_row = int(table.ref.split(":")[0][1:])
    existing_names = [col.name for col in table.tableColumns]

    # Erste Spalte der Tabelle bestimmen (z.B. 'B2:Y92' -> Spalte B)
    first_col_letter = "".join(ch for ch in table.ref.split(":")[0] if ch.isalpha())
    first_col_idx = 0
    for ch in first_col_letter:
        first_col_idx = first_col_idx * 26 + (ord(ch.upper()) - ord("A") + 1)

    table_end_col_idx = first_col_idx + len(existing_names) - 1

    # Lose Inhalte jenseits der Tabelle (z.B. Notizen) nicht überschreiben.
    # Da eine Excel-Tabelle zusammenhängend ist, werden solche Spalten mit
    # einem generischen Header in die Tabelle aufgenommen (bleiben erhalten,
    # werden vom Import ignoriert); neue Spalten beginnen dahinter.
    stray_columns: list[tuple[int, str]] = []  # (Spaltenindex, Header-Name)
    for col_idx in range(table_end_col_idx + 1, sheet.max_column + 1):
        filled_rows = [
            row
            for row in range(1, sheet.max_row + 1)
            if row != header_row and sheet.cell(row=row, column=col_idx).value is not None
        ]
        if not filled_rows:
            continue
        header_value = sheet.cell(row=header_row, column=col_idx).value
        if header_value is None or str(header_value).strip() == "":
            header_value = f"frei_{get_column_letter(col_idx)}"
            sheet.cell(row=header_row, column=col_idx, value=header_value)
        logger.warning(
            "Inhalte ausserhalb der Tabelle '{}' in Spalte {} (Zeilen {}-{}) gefunden — "
            "werden als Spalte '{}' in die Tabelle aufgenommen (Import ignoriert sie). "
            "Bitte aufräumen.",
            CLIENT_TABLE_NAME,
            get_column_letter(col_idx),
            filled_rows[0],
            filled_rows[-1],
            header_value,
        )
        stray_columns.append((col_idx, str(header_value)))

    next_col_idx = table_end_col_idx + 1
    if stray_columns:
        next_col_idx = stray_columns[-1][0] + 1

    added: list[str] = []
    for field in ACCORDIX_CLIENT_FIELDS:
        if field in existing_names:
            continue
        sheet.cell(row=header_row, column=next_col_idx, value=field)
        added.append(field)
        next_col_idx += 1

    if not added and not stray_columns:
        logger.info("Alle Accordix-Spalten sind in '{}' bereits vorhanden.", CLIENT_TABLE_NAME)
    elif added or stray_columns:
        # Tabellen-Referenz und Spaltendefinitionen erweitern
        last_row = int(table.ref.split(":")[1][1:])
        new_ref = (
            f"{get_column_letter(first_col_idx)}{header_row}:"
            f"{get_column_letter(next_col_idx - 1)}{last_row}"
        )
        table.ref = new_ref
        max_id = max(col.id for col in table.tableColumns)
        # Physikalische Reihenfolge: bisherige Spalten, dann Fund-Spalten, dann neue
        for offset, (_, name) in enumerate(stray_columns, start=1):
            table.tableColumns.append(TableColumn(id=max_id + offset, name=name))
        for offset, field in enumerate(added, start=len(stray_columns) + 1):
            table.tableColumns.append(TableColumn(id=max_id + offset, name=field))
        if added:
            logger.info("Tabelle '{}' um Spalten erweitert: {}", CLIENT_TABLE_NAME, ", ".join(added))

    # Datenvalidierungen (Dropdowns) und Formate auf die neuen Spalten
    all_names = [col.name for col in table.tableColumns]
    for field in ACCORDIX_CLIENT_FIELDS:
        col_idx = first_col_idx + all_names.index(field)
        col_letter = get_column_letter(col_idx)
        cell_range = f"{col_letter}{_FIRST_DATA_ROW}:{col_letter}{_LAST_VALIDATED_ROW}"

        if field in FIELD_VALUE_LISTS:
            # Bestehende Validierung für diesen Bereich entfernen (Idempotenz)
            name = f"{_PREFIX}{field}"
            sheet.data_validations.dataValidation = [
                dv for dv in sheet.data_validations.dataValidation if str(dv.sqref) != cell_range
            ]
            dv = DataValidation(type="list", formula1=name, allow_blank=True, showDropDown=False)
            dv.error = "Wert nicht in der Accordix-Werteliste (siehe Blatt 'Wertelisten')."
            dv.errorTitle = "Ungültiger Wert"
            sheet.add_data_validation(dv)
            dv.add(cell_range)
        elif field == DATE_COLUMN:
            for row in range(_FIRST_DATA_ROW, _LAST_VALIDATED_ROW + 1):
                sheet[f"{col_letter}{row}"].number_format = DATE_FORMAT

        current_width = sheet.column_dimensions[col_letter].width or 0
        sheet.column_dimensions[col_letter].width = max(current_width, 14)


def extend_masterdata_file(excel_path: Path) -> Path:
    """
    Erweitert eine Stammdaten-Datei um die Accordix-Struktur (idempotent).

    Args:
        excel_path: Pfad zur Excel-Datei (wird in-place angepasst).

    Returns:
        Pfad zur angepassten Datei.
    """
    if not excel_path.exists():
        raise FileNotFoundError(f"Stammdaten-Datei nicht gefunden: {excel_path}")

    workbook = load_workbook(excel_path)
    _ensure_valuelist_sheet(workbook)
    _ensure_client_columns(workbook)
    workbook.save(excel_path)
    logger.info("Stammdaten-Datei erweitert: {}", excel_path)
    return excel_path

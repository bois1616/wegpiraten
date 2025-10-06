"""
Importiert Stammdaten aus einer bestehenden Excel-Datei (mit mehreren Tabellen/Excel-Tabellen) in die Projekt-SQLite-DB.
Die Ziel-DB wird automatisch im local_data_path unterhalb des Projekt-Roots angelegt (Pfad und Name aus Config).
Die Quelldatei (Excel) wird aus shared_data_path geladen.
Verwendet zentrale Config, Entity-Modelle aus shared_modules.entity_config und Feld-Mappings aus der Config.
"""

from pathlib import Path
import sqlite3
from typing import Dict, Type, Any
import pandas as pd
from loguru import logger
import openpyxl

from shared_modules.config import Config
from shared_modules.entity_config import Employee, Client, Payer, ServiceRequester
from shared_modules.mapping_entry import MappingEntry

def sql_type(py_type: str) -> str:
    # Mapping von Python-Typnamen (als String) auf SQLite-Typen
    return {
        "str": "TEXT",
        "float": "REAL",
        "int": "INTEGER",
        "bool": "INTEGER"
    }.get(py_type, "TEXT")

def create_target_tables(conn, table_mappings, field_mappings):
    cur = conn.cursor()
    for excel_table, table_cfg in table_mappings.items():
        target_table = table_cfg["target"]
        entity = table_cfg["entity"]
        primary_key = table_cfg.get("primary_key")
        fields = field_mappings[entity]
        columns = []
        for excel_col, entry in fields.items():
            col_name = entry.field
            col_type = sql_type(entry.type)
            if col_name == primary_key:
                columns.append(f"{col_name} {col_type} PRIMARY KEY")
            else:
                columns.append(f"{col_name} {col_type}")
        columns_sql = ",\n    ".join(columns)
        sql = f"CREATE TABLE IF NOT EXISTS {target_table} (\n    {columns_sql}\n)"
        cur.execute(sql)
    conn.commit()

def get_type_from_str(type_str: str):
    """
    Wandelt einen Typnamen als String in einen Python-Typ um.
    """
    mapping = {
        "str": str,
        "float": float,
        "int": int,
        "bool": bool,
    }
    return mapping.get(type_str, str)

def map_row(row: pd.Series, mapping: Dict[str, MappingEntry], required_fields: list) -> Dict[str, Any]:
    """
    Mappt die Felder einer Zeile gemäß dem Mapping-Dict aus der Config.
    Führt erforderliche Typkonvertierungen durch und ergänzt fehlende Felder mit None.
    """
    result = {}
    for excel_col, entry in mapping.items():
        field_name = entry.field
        field_type = get_type_from_str(entry.type)
        value = row.get(excel_col)
        if value is not None:
            try:
                # Spezialfall: float auf int, falls nötig
                if field_type is int and isinstance(value, float) and value.is_integer():
                    value = int(value)
                else:
                    value = field_type(value)
            except Exception:
                logger.warning(f"Typkonvertierung für Feld '{field_name}' fehlgeschlagen, Wert: {value}")
        result[field_name] = value
    # Fehlende Pflichtfelder ergänzen
    for field in required_fields:
        if field not in result:
            result[field] = None
    return result

def read_excel_table(file_path: Path, table_name: str) -> pd.DataFrame:
    """
    Liest eine benannte Tabelle (Excel Table, nicht Sheet!) aus einer Excel-Datei als DataFrame.
    """
    wb = openpyxl.load_workbook(file_path, data_only=True)
    for ws in wb.worksheets:
        if table_name in ws.tables:
            table = ws.tables[table_name]
            ref = table.ref  # z.B. 'A1:F20'
            min_col, min_row, max_col, max_row = openpyxl.utils.range_boundaries(ref)
            data = []
            for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                                    min_col=min_col, max_col=max_col):
                data.append([cell.value for cell in row])
            df = pd.DataFrame(data[1:], columns=data[0])  # Erste Zeile als Header
            return df
    raise ValueError(f"Tabelle {table_name} nicht gefunden.")

def import_table(
    source_excel: Path,
    target_conn: sqlite3.Connection,
    table_name: str,
    target_table: str,
    model: Type,
    mapping: Dict[str, MappingEntry],
) -> None:
    """
    Liest alle Daten aus der angegebenen Excel-Tabelle (nicht Sheet!), mappt die Felder und schreibt sie in die Zieltabelle.
    """
    logger.info(f"Importiere Excel-Tabelle {table_name} → {target_table}")
    try:
        excel_table = read_excel_table(source_excel, table_name)
    except Exception as e:
        logger.error(f"Fehler beim Lesen der Excel-Tabelle {table_name}: {e}")
        return

    required_fields = list(model.model_fields.keys())
    records = []
    for _, row in excel_table.iterrows():
        if row.isnull().all():
            continue
        try:
            mapped = map_row(row, mapping, required_fields)
            rec = model(**mapped)
            records.append(rec.model_dump())
        except Exception as e:
            logger.error(f"Fehler beim Validieren eines Datensatzes aus {table_name}: {e}")

    if records:
        df_target = pd.DataFrame(records)
        try:
            df_target.to_sql(target_table, target_conn, if_exists="append", index=False)
            logger.success(f"{len(df_target)} Datensätze in {target_table} importiert.")
        except Exception as e:
            logger.error(f"Fehler beim Schreiben in {target_table}: {e}")
    else:
        logger.warning(f"Keine gültigen Datensätze für {target_table} gefunden.")

def main() -> None:
    """
    Hauptfunktion: Liest die Konfiguration, legt die Ziel-DB an (falls erforderlich),
    legt die Tabellen an und importiert die Stammdaten.
    """
    # Lade zentrale Konfiguration
    config_path = Path(__file__).parent.parent.parent / ".config" / "wegpiraten_config.yaml"
    config = Config()
    config.load(config_path)

    # Ermittle Pfade aus der Config
    prj_root = Path(config.get_structure().prj_root)
    shared_data_path = getattr(config.get_structure(), "shared_data_path", "shared_data")
    local_data_path = getattr(config.get_structure(), "local_data_path", "data")
    sqlite_db_name = config.get("sqlite_db_name")
    source_db_name = config.get("db_name")

    # Ziel-DB: im local_data_path unterhalb von prj_root
    target_db_path = prj_root / local_data_path / sqlite_db_name
    target_db_path.parent.mkdir(parents=True, exist_ok=True)

    # Quelle: Excel-Datei im shared_data_path unterhalb von prj_root
    source_excel_path = prj_root / shared_data_path / source_db_name

    logger.info(f"Quelle: {source_excel_path}")
    logger.info(f"Ziel:   {target_db_path}")

    if not source_excel_path.exists():
        logger.error(f"Quelldatei nicht gefunden: {source_excel_path}")
        return

    # Lade Tabellen-Mapping aus der Config (z.B. aus mappings_config)
    # Beispielstruktur in der Config:
    # table_mappings:
    #   MD_MA:
    #     entity: employee
    #     target: employees
    #   ...
    table_mappings = config.get("table_mappings")
    field_mappings = config.get("field_mappings")

    # Entity-Name zu Modell
    entity_models = {
        "employee": Employee,
        "service_requester": ServiceRequester,
        "payer": Payer,
        "client": Client,
    }

    # Generiere TABLES dynamisch aus der Config
    TABLES = {}
    for excel_table, table_cfg in table_mappings.items():
        entity = table_cfg["entity"]
        TABLES[excel_table] = {
            "target": table_cfg["target"],
            "model": entity_models[entity],
            "mapping": field_mappings[entity],
        }

    with sqlite3.connect(target_db_path) as target_conn:
        create_target_tables(target_conn, table_mappings, field_mappings)
        for source_table, meta in TABLES.items():
            try:
                import_table(
                    source_excel_path,
                    target_conn,
                    source_table,      # Tabellenname in Excel
                    meta["target"],
                    meta["model"],
                    meta["mapping"],
                )
            except Exception as e:
                logger.error(f"Fehler beim Import von {source_table}: {e}")
    logger.info("Stammdaten-Import abgeschlossen.")

if __name__ == "__main__":
    main()
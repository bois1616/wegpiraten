"""Integration des neuen Schemas ohne Kundendaten."""

import json
import sqlite3
from datetime import date, datetime
from pathlib import Path

import pytest
import yaml
from openpyxl import Workbook, load_workbook
from pytest import MonkeyPatch

from data_imports.batch_import_timesheets import TimeSheetsImporter
from data_imports.import_masterdata import run_import
from data_imports.legacy_mandates import LegacyMandateResolver
from invoices.modules.invoice_filter import InvoiceFilter
from invoices.modules.invoice_processor import InvoiceProcessor
from shared_modules.config import Config
from shared_modules.month_period import get_month_period
from tests.fixtures.create_masterdata import create_workbook
from time_sheets.modules.client_data import load_active_client_headers, load_internal_timesheet_headers


@pytest.fixture
def imported(tmp_path: Path, monkeypatch: MonkeyPatch) -> Config:
    """Isolierte Config, synthetisches Workbook und eigene DB je Test."""
    root = Path(__file__).resolve().parents[1]
    data = yaml.safe_load((root / ".config/wegpiraten_config.yaml").read_text())
    data["structure"]["prj_root"] = str(tmp_path)
    data["structure"]["template_path"] = str(root / "templates")
    data["database"]["legacy_mandate_mapping"] = "legacy.json"
    data["logging"] = {"log_level": "WARNING", "log_file": str(tmp_path / "test.log")}
    for directory in ("data", "import", "done", "output", ".tmp", ".logs", "graphics"):
        (tmp_path / directory).mkdir()
    path = tmp_path / "config.yaml"
    path.write_text(yaml.safe_dump(data))
    (tmp_path / "legacy.json").write_text(json.dumps({"C9001": ["A26091", "A26092"]}))
    monkeypatch.setattr(Config, "_instance", None)
    config = Config(path)
    run_import(config, create_workbook(tmp_path / "master.xlsx"))
    return config


def test_view_index_budget_sentinel_and_timesheet_headers(imported: Config) -> None:
    """Indexkind kann unbetreut sein; SA hat keine künstliche Person oder Betreuung."""
    with sqlite3.connect(imported.get_db_path()) as conn:
        assert conn.execute(
            "SELECT index_person_id,allowed_travel_time FROM v_mandate WHERE mandate_id='A26093'"
        ).fetchone() == ("C9002", 30)
        assert conn.execute("SELECT count(*) FROM v_mandate").fetchone()[0] == 5
        assert conn.execute("SELECT count(*) FROM mandate_person WHERE mandate_id='SA'").fetchone()[0] == 0
        assert not conn.execute("SELECT 1 FROM sqlite_master WHERE name='clients'").fetchone()
    headers = load_active_client_headers(imported.get_db_path(), "2026-09")
    assert {h.client_id for h in headers} == {"A26091", "A26093"}
    assert len(load_internal_timesheet_headers(imported.get_db_path())) == 2


def test_legacy_renewal_dates_overlap_and_unknown(imported: Config) -> None:
    resolver = LegacyMandateResolver(imported.get_db_path(), Path(imported.structure.prj_root) / "legacy.json")
    assert resolver.resolve("C9001", date(2026, 9, 30)) == ("A26091", False)
    assert resolver.resolve("C9001", date(2026, 10, 1)) == ("A26092", False)
    assert resolver.resolve("SA", date(2026, 9, 1)) == ("SA", False)
    with pytest.raises(ValueError):
        resolver.resolve("C9999", date(2026, 9, 1))
    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.execute("UPDATE mandate SET start_date='2026-09-01' WHERE mandate_id='A26092'")
    with pytest.raises(ValueError):
        resolver.resolve("C9001", date(2026, 9, 1))


def test_timesheet_crosses_renewal_boundary_and_file_failure_is_atomic(imported: Config) -> None:
    importer = TimeSheetsImporter(imported)
    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.execute("UPDATE mandate SET allowed_direct_effort=180 WHERE mandate_id='A26092'")
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["F5"] = "E9001"
    sheet["F8"] = "C9001"
    sheet["C6"] = datetime(2026, 9, 1)
    sheet["A10"] = "Uhrzeit"
    sheet["B10"] = "Datum"
    sheet["C10"] = "Fahrzeit"
    sheet["D10"] = "Direkter Fallkontakt"
    sheet["E10"] = "Indirekte Fallbearbeitung"
    sheet["H10"] = "Notizen"
    sheet["B11"] = datetime(2026, 9, 30)
    sheet["D11"] = 15
    sheet["H11"] = "Ohne Berechnung"
    sheet["B12"] = datetime(2026, 10, 1)
    sheet["D12"] = 30
    path = importer.source_dir / "crossing.xlsx"
    workbook.save(path)
    count, _, _, _ = importer.process_file(path, "2026-09", get_month_period("2026-09"))
    assert count == 2
    with sqlite3.connect(imported.get_db_path()) as conn:
        rows = conn.execute("SELECT mandate_id,direct_time,notes FROM service_data ORDER BY service_date").fetchall()
    assert rows == [("A26091", 15, "Ohne Berechnung"), ("A26092", 30, None)]
    with sqlite3.connect(imported.get_db_path()) as conn:
        assert conn.execute("SELECT allowed_direct_effort FROM service_data ORDER BY service_date").fetchall() == [
            (600,),
            (180,),
        ]
    processor = InvoiceProcessor(imported, InvoiceFilter(invoice_month="09.2026", client_list=["C9002"]))
    data = processor._load_service_data(get_month_period("2026-09"))
    assert len(data) == 1 and data.iloc[0]["client_id"] == "A26091"
    assert data.iloc[0]["hourly_rate"] == 100
    # Die zweite Datei enthält eine unauflösbare Zeile: auch die erste darf nicht gespeichert werden.
    sheet["B12"] = datetime(2025, 1, 1)
    bad = importer.source_dir / "bad.xlsx"
    workbook.save(bad)
    count, _, _, _ = importer.process_file(bad, "2026-09", get_month_period("2026-09"))
    assert count == 0 and bad.exists()
    with sqlite3.connect(imported.get_db_path()) as conn:
        assert conn.execute("SELECT count(*) FROM service_data").fetchone()[0] == 2


def test_twins_use_workbook_order_and_missing_reference_stays_visible(imported: Config) -> None:
    """Reihenfolge bleibt erhalten; fehlerhafte Verweise verschwinden nicht aus der View."""
    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.execute("UPDATE person SET date_of_birth='2018-01-01' WHERE person_id='C9001'")
        conn.execute("UPDATE person SET source_row=1 WHERE person_id='C9002'")
        assert conn.execute("SELECT index_person_id FROM v_mandate WHERE mandate_id='A26091'").fetchone()[0] == "C9002"
        conn.execute("UPDATE mandate SET contact_person_id='AP_MISSING' WHERE mandate_id='A26091'")
        assert conn.execute("SELECT count(*) FROM v_mandate WHERE mandate_id='A26091'").fetchone()[0] == 1
        assert conn.execute("SELECT sr_ap_last_name FROM v_mandate WHERE mandate_id='A26091'").fetchone()[0] is None


def test_accordix_care_continuation_and_exit_are_not_authorisation_end(imported: Config) -> None:
    """Zwei passende Folgeaufträge melden dieselbe Betreuung einmal; Ende ist echter Austritt."""
    from reports.accordix_report import create_accordix_report

    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.execute("UPDATE mandate SET start_date='2026-09-15' WHERE mandate_id='A26092'")
    output = create_accordix_report(imported, "2026-09")
    sheet = load_workbook(output, data_only=True)["Ambulant"]
    populated = [r for r in sheet.iter_rows(min_row=7, values_only=True) if r[0]]
    assert len(populated) == 3
    assert all(r[18] is None for r in populated)
    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.execute(
            "UPDATE mandate_person SET end_date='2026-09-20', leaving_reason='Anderer', custom_leaving_reason='Testende' WHERE mandate_id='A26092'"
        )
    output = create_accordix_report(imported, "2026-09")
    sheet = load_workbook(output, data_only=True)["Ambulant"]
    populated = [r for r in sheet.iter_rows(min_row=7, values_only=True) if r[0]]
    ended = [r for r in populated if r[18]]
    assert len(ended) == 2
    assert all(r[18] == "20.09.2026" and r[20] == "Anderer" and r[21] == "Testende" for r in ended)


def test_invoice_keeps_zero_charge_position_and_skips_unusable_tariff(
    imported: Config, monkeypatch: MonkeyPatch
) -> None:
    """Kostenfreie Minuten sind sichtbar, erhöhen den Betrag aber nicht."""
    from PyPDF2 import PdfWriter

    from invoices.modules.document_utils import DocumentUtils
    from invoices.modules.invoice_context import InvoiceContext

    TimeSheetsImporter(imported)
    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.executemany(
            """INSERT INTO service_data (mandate_id,employee_id,service_date,service_type,
            direct_time,notes,reporting_month,allowed_direct_effort) VALUES ('A26091','E9001','2026-09-01','SPF',?,?,'2026-09',600)""",
            [(15, "Ohne Berechnung"), (45, None)],
        )
    contexts: list[InvoiceContext] = []

    def fake_pdf(docx_path: Path, pdf_path: Path, context: InvoiceContext) -> Path:
        """PDF-Konverter isolieren; die eigentliche DOCX-Erzeugung läuft unverändert."""
        contexts.append(context)
        writer = PdfWriter()
        writer.add_blank_page(width=595, height=842)
        with pdf_path.open("wb") as handle:
            writer.write(handle)
        return pdf_path

    monkeypatch.setattr(DocumentUtils, "docx_to_pdf", staticmethod(fake_pdf))
    InvoiceProcessor(imported, InvoiceFilter(invoice_month="09.2026")).run()
    assert len(contexts) == 1
    positions = contexts[0].data["positions"]
    assert len(positions) == 2
    assert [p["Kosten"] for p in positions] == [0, 75]
    assert contexts[0].data["summe_kosten"] == 75
    assert contexts[0].data["client"].social_security_number == ""
    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.execute("UPDATE service_types SET hourly_rate=NULL WHERE service_type_id='ST01'")
    contexts.clear()
    InvoiceProcessor(imported, InvoiceFilter(invoice_month="09.2026")).run()
    assert not contexts


def test_new_timesheet_can_be_filled_and_imported(imported: Config, monkeypatch: MonkeyPatch) -> None:
    """Erzeugung und Import teilen dieselben konfigurierten Kopfzellen und Minutenbudgets."""
    from time_sheets.modules.time_sheet_factory import TimeSheetFactory

    monkeypatch.setenv("SHEET_PASSWORD", "synthetic-test-only")
    factory = TimeSheetFactory(imported)
    headers = factory.fetch_reporting_data("2026-09")
    header = next(h for h in headers if h.client_id == "A26091" and h.employee_id == "E9001")
    path = factory.create_reporting_sheet(header, datetime(2026, 9, 1), output_path=imported.get_imports_path())
    book = load_workbook(path)
    sheet = book.active
    assert sheet is not None
    assert sheet["F8"].value == "A26091"
    sheet["B11"] = datetime(2026, 9, 1)
    sheet["D11"] = 30
    book.save(path)
    importer = TimeSheetsImporter(imported)
    count, _, _, _ = importer.process_file(path, "2026-09", get_month_period("2026-09"))
    assert count == 1
    with sqlite3.connect(imported.get_db_path()) as conn:
        assert conn.execute("SELECT mandate_id,direct_time,allowed_travel_time FROM service_data").fetchone() == (
            "A26091",
            30,
            30,
        )


def test_terminlisten_month_status_assignment_and_excel(imported: Config) -> None:
    """Monatsgrenzen, Status und P-Zuordnung; auch MA ohne TS erhalten eine Datei."""
    from reports.terminliste import create_terminlisten

    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.executemany(
            "INSERT INTO report(report_id,mandate_id,due_date,report_form,status,notes) VALUES (?,?,?,?,?,?)",
            [
                ("R9002", "A26093", "2026-10-01", "Zwischenbericht", None, "=Freitext"),
                ("R9003", "A26094", "2026-10-31", "Abschlussbericht", "offen", "Nach Auftragsende"),
                ("R9004", "A26092", "2026-09-30", "Bericht", "offen", None),
                ("R9005", "A26092", "2026-11-01", "Bericht", "offen", None),
                ("R9006", "A26092", "2026-10-20", "Bericht", "erledigt", None),
                ("R9007", "A26092", "2026-10-20", "Bericht", "entfällt", None),
                ("R9008", "A26092", None, "Bericht", "offen", None),
            ],
        )
    files = create_terminlisten(imported, "10.2026")
    assert len(files) == 3
    for path in files:
        workbook = load_workbook(path)
        sheet = workbook["Termine"]
        assert sheet.freeze_panes == "C5"
        assert sheet.page_setup.fitToWidth == 1
        if "E9001" in path.name:
            assert [sheet.cell(row, 9).value for row in (5, 6)] == ["R9001", "R9003"]
            assert sheet["A6"].value == datetime(2026, 10, 31)
            assert sheet["A6"].number_format == "dd.mm.yyyy"
        elif "E9002" in path.name:
            assert sheet["I5"].value == "R9002"
            assert sheet["H5"].value == "=Freitext"
            assert sheet["H5"].data_type == "s"
        else:
            assert sheet["A5"].value == "Keine fälligen Berichte."
        workbook.close()
    # Gleicher Monat überschreibt die Dateien statt weitere anzulegen.
    assert create_terminlisten(imported, "2026-10") == files


def test_terminlisten_unassignable_reports_warn_and_continue(imported: Config) -> None:
    """Fehlende oder mehrfache P-Zuordnung blockiert die übrigen Dateien nicht."""
    from loguru import logger

    from reports.terminliste import create_terminlisten

    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.execute("DELETE FROM relation_mandate_emp WHERE mandate_id='A26092'")
        conn.execute("INSERT INTO relation_mandate_emp(mandate_id,employee_id,role) VALUES ('A26093','E9001','P')")
        conn.execute("INSERT INTO relation_mandate_emp(mandate_id,employee_id,role) VALUES ('A26094','E_MISSING','P')")
        conn.execute("DELETE FROM relation_mandate_emp WHERE mandate_id='A26094' AND employee_id='E9001'")
        conn.executemany(
            "INSERT INTO report(report_id,mandate_id,due_date,status) VALUES (?,?,?,'offen')",
            [
                ("R9002", "A26093", "2026-10-15"),
                ("R9003", "A26094", "2026-10-15"),
                ("R9004", "A_MISSING", "2026-10-15"),
                ("R9005", "A26091", "2026-10-15"),
            ],
        )
    messages: list[str] = []
    sink = logger.add(lambda message: messages.append(str(message)), level="WARNING")
    try:
        files = create_terminlisten(imported, "2026-10")
    finally:
        logger.remove(sink)
    assert len(files) == 3
    for path in files:
        workbook = load_workbook(path)
        if "E9001" in path.name:
            assert workbook["Termine"]["I5"].value == "R9005"
        else:
            assert workbook["Termine"]["A5"].value == "Keine fälligen Berichte."
        workbook.close()
    assert len(messages) == 4
    assert any("R9001 / Auftrag A26092" in message and "keine Person mit Rolle P" in message for message in messages)
    assert any("R9002" in message and "mehrere Personen mit Rolle P" in message for message in messages)
    assert any("R9003" in message and "Mitarbeiter-Stammdaten" in message for message in messages)
    assert any("R9004" in message and "Auftrag fehlt" in message for message in messages)


def test_terminliste_cli(imported: Config) -> None:
    """CLI-Aufruf erstellt Dateien und meldet einen ungültigen Monat als Fehler."""
    from typer.testing import CliRunner

    from cli import app

    runner = CliRunner()
    config_path = str(Path(imported.structure.prj_root) / "config.yaml")
    result = runner.invoke(app, ["terminliste", "2026-10", "--config", config_path])
    assert result.exit_code == 0, result.output
    assert "3 Terminlisten erstellt" in result.output
    assert len(list((imported.get_output_path() / "Terminlisten_2026-10").glob("*.xlsx"))) == 3
    result = runner.invoke(app, ["terminliste", "2026-13", "--config", config_path])
    assert result.exit_code == 1


def test_report_history_import_and_due_day_assignment(imported: Config) -> None:
    """Datierte XLSX-Zuordnung überlebt den Neuaufbau und gilt für beide Reports."""
    from openpyxl.worksheet.table import Table

    from reports.klientenuebersicht import create_klientenuebersicht
    from reports.report_data import load_report_data
    from reports.terminliste import create_terminlisten

    source = create_workbook(Path(imported.structure.prj_root) / "with_history.xlsx")
    workbook = load_workbook(source)
    sheet = workbook.create_sheet("Zuständigkeitsverlauf")
    sheet.append(["assignment_id", "mandate_id", "employee_id", "role", "valid_from", "valid_to", "notes"])
    sheet.append(["Z1", "A26092", "E9001", "P", datetime(2026, 1, 1), datetime(2026, 7, 31), None])
    sheet.append(["Z2", "A26092", "E9002", "P", datetime(2026, 8, 1), datetime(2026, 12, 31), None])
    sheet.append(["Z3", "A26092", "E9003", "P", datetime(2027, 1, 1), None, None])
    sheet.append(["Z4", "A26092", "E9001", "S", datetime(2026, 8, 1), None, None])
    sheet.add_table(Table(displayName="assignment_history", ref="A1:G5"))
    workbook.save(source)
    workbook.close()
    # Ausschliesslich die isolierte Test-DB verwerfen, niemals die Kunden-DB.
    imported.get_db_path().unlink()
    run_import(imported, source)
    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.executemany(
            "INSERT INTO report(report_id,mandate_id,due_date,report_form,status) VALUES (?,?,?,'Bericht','offen')",
            [
                ("R9010", "A26092", "2026-07-31"),
                ("R9011", "A26092", "2026-08-01"),
                ("R9012", "A26092", "2026-12-31"),
                ("R9013", "A26092", "2027-01-01"),
            ],
        )
        _, reports = load_report_data(conn)
        assert {r.report_id: r.employee_id for r in reports} == {
            "R9001": "E9002",
            "R9010": "E9001",
            "R9011": "E9002",
            "R9012": "E9002",
            "R9013": "E9003",
        }
        assert all("Fälligkeitstag" in r.assignment_basis for r in reports)
        count_before = conn.execute("SELECT count(*) FROM report").fetchone()[0]
    files = create_terminlisten(imported, "2026-10")
    own = next(path for path in files if "E9002" in path.name)
    workbook = load_workbook(own)
    assert workbook["Termine"]["I5"].value == "R9001"
    workbook.close()
    overview = load_workbook(create_klientenuebersicht(imported, "2026-10"))
    assert "Übersicht" not in overview.sheetnames
    sheets = [sheet for sheet in overview if sheet.title.startswith("E900")]
    assigned = next(sheet for sheet in sheets if "E9002" in sheet["A2"].value)
    assert "R9001" in assigned["X6"].value
    successor = next(sheet for sheet in sheets if "E9003" in sheet["A2"].value)
    assert "R9013" in successor["AA6"].value
    overview.close()
    with sqlite3.connect(imported.get_db_path()) as conn:
        assert conn.execute("SELECT count(*) FROM report").fetchone()[0] == count_before
    source_workbook = load_workbook(source)
    assert source_workbook["Zuständigkeitsverlauf"].max_row == 5
    source_workbook.close()


def test_old_contact_columns_and_no_history_remain_importable(imported: Config) -> None:
    """Neue Kontaktfelder und Historientabelle sind bei alten Mappen optional."""
    source = create_workbook(Path(imported.structure.prj_root) / "old_columns.xlsx")
    workbook = load_workbook(source)
    sheet = workbook["Ansprechpersonen"]
    sheet.delete_cols(7, 2)
    sheet.tables["masterdata_contact_person"].ref = "A2:F3"
    workbook.save(source)
    workbook.close()
    imported.get_db_path().unlink()
    run_import(imported, source)
    with sqlite3.connect(imported.get_db_path()) as conn:
        assert conn.execute("SELECT phone,email FROM masterdata_contact_person").fetchone() == (None, None)
        assert conn.execute("SELECT count(*) FROM assignment_history").fetchone()[0] == 0


@pytest.mark.parametrize("problem", ["gap", "overlap", "invalid", "missing_employee"])
def test_report_history_never_falls_back_for_bad_period(imported: Config, problem: str) -> None:
    """Eine lückenhafte/mehrdeutige Historie darf die aktuelle P-Zuordnung nicht nutzen."""
    from reports.klientenuebersicht import create_klientenuebersicht
    from reports.report_data import assignment_problem, load_report_data

    with sqlite3.connect(imported.get_db_path()) as conn:
        start = "2026-11-01" if problem == "gap" else "bad date" if problem == "invalid" else "2026-01-01"
        employee = "E_MISSING" if problem == "missing_employee" else "E9002"
        conn.execute(
            "INSERT INTO assignment_history(assignment_id,mandate_id,employee_id,role,valid_from) VALUES ('Z1','A26092',?,'P',?)",
            (employee, start),
        )
        if problem == "overlap":
            # Auch zwei Perioden derselben Person sind eine Überschneidung.
            conn.execute(
                "INSERT INTO assignment_history(assignment_id,mandate_id,employee_id,role,valid_from) VALUES ('Z2','A26092','E9002','P','2026-10-01')"
            )
        employees, reports = load_report_data(conn)
        assert assignment_problem(reports[0], {e.emp_id for e in employees}) is not None
        assert reports[0].employee_id != "E9001"
    workbook = load_workbook(create_klientenuebersicht(imported, "2026-10"))
    assert workbook["Nicht zugeordnet"]["A2"].value == "R9001"
    assert workbook["Berichtsdetails"]["A2"].value == "R9001"
    workbook.close()


def test_klientenuebersicht_grid_status_details_and_cli(imported: Config) -> None:
    """Gemischte Termine, Folgejahr und fehlende Daten bleiben ohne Kopien sichtbar."""
    from typer.testing import CliRunner

    from cli import app
    from reports.klientenuebersicht import create_klientenuebersicht

    with sqlite3.connect(imported.get_db_path()) as conn:
        conn.executemany(
            "INSERT INTO report(report_id,mandate_id,due_date,report_form,status,notes) VALUES (?,?,?,?,?,?)",
            [
                ("R9010", "A26092", "2026-10-20", "Zwischenbericht", "erledigt", "=Kein Formeltext"),
                ("R9011", "A26092", "2027-01-01", "Abschlussbericht", "entfällt", None),
                ("R9012", "A26092", None, "Bericht", "offen", None),
                ("R9013", "A26092", "2028-01-01", "Bericht", "offen", None),
                ("R9014", "A_MISSING", "2026-10-15", "Bericht", None, None),
            ],
        )
    path = create_klientenuebersicht(imported, "2026-10")
    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Legende",
        "E9001 Test",
        "E9003 Test",
        "E9002 Test",
        "Berichtsdetails",
        "Nicht zugeordnet",
    ]
    sheet = workbook["E9001 Test"]
    assert sheet["A1"].value is None
    assert sheet["O4"].value == "Jan" and sheet["AA4"].value == "Jan"
    assert sheet["X6"].fill.fgColor.rgb == "FFE2E8F0"
    assert "R9001" in sheet["X6"].value and "R9010" in sheet["X6"].value
    assert "erledigt" in sheet["X6"].value
    assert "R9011" in sheet["AA6"].value and "entfällt" in sheet["AA6"].value
    assert sheet["AA6"].fill.fgColor.rgb == "FF00B0F0"
    assert sheet["J6"].value == 0.5 and sheet["M6"].value == 15.5
    assert sheet.freeze_panes == "O6" and sheet.max_column == 38
    details = workbook["Berichtsdetails"]
    assert details.max_row == 7
    assert sorted(details.cell(row, 1).value for row in range(2, 8)) == [
        "R9001",
        "R9010",
        "R9011",
        "R9012",
        "R9013",
        "R9014",
    ]
    text = next(details.cell(row, 10) for row in range(2, 8) if details.cell(row, 1).value == "R9010")
    assert text.value == "=Kein Formeltext" and text.data_type == "s"
    assert all(
        "historische Zuständigkeit unbekannt" in details.cell(row, 8).value
        for row in range(2, 8)
        if details.cell(row, 7).value
    )
    assert workbook["Nicht zugeordnet"].max_row == 4
    workbook.close()
    assert create_klientenuebersicht(imported, "2026-10") == path
    runner = CliRunner()
    config_path = str(Path(imported.structure.prj_root) / "config.yaml")
    assert runner.invoke(app, ["klientenuebersicht", "2026-10", "--config", config_path]).exit_code == 0
    assert runner.invoke(app, ["klientenuebersicht", "2026-13", "--config", config_path]).exit_code == 1

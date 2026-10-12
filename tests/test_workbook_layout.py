"""Spaltenwechsel erhalten Eingaben und Bezüge auch auf anderen Blättern."""

from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from data_imports import workbook_layout
from data_imports.workbook_layout import group_input_columns, remap_formula, remap_reference
from data_imports.workbook_package import NS, WorkbookPackage, tag


def test_column_grouping_preserves_values_and_dependencies(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Ansprechpersonen"
    sheet.append(["Beschreibung"])
    sheet.append(["Nr", "Name", "Anzeige", "Telefon"])
    sheet.append(["id", "name", "display", "phone"])
    sheet.append(["AP1", "Name", '=B4&" Telefon D4"', "+4101234"])
    sheet.add_table(Table(displayName="contacts", ref="A3:D4"))
    validation = DataValidation(type="list", formula1="=$D$4:$D$5")
    sheet.add_data_validation(validation)
    validation.add("D4:D8")
    sheet.conditional_formatting.add("C4:C8", FormulaRule(formula=['D4<>""']))
    workbook.create_sheet("Andere")["A1"] = "='Ansprechpersonen'!$D$4"
    workbook.defined_names.add(DefinedName("phones", attr_text="'Ansprechpersonen'!$D$4:$D$8"))
    workbook.save(source)
    monkeypatch.setattr(workbook_layout, "_TABLES", {"Ansprechpersonen": "contacts"})
    package = WorkbookPackage(source)
    mappings = group_input_columns(package)
    assert mappings["Ansprechpersonen"] == {1: 1, 2: 2, 4: 3, 3: 4}
    destination = tmp_path / "grouped.xlsx"
    package.write(destination)
    result = load_workbook(destination)
    assert result["Ansprechpersonen"]["C4"].value == "+4101234"
    assert result["Ansprechpersonen"]["D4"].value == '=B4&" Telefon D4"'
    assert result["Andere"]["A1"].value == "='Ansprechpersonen'!$C$4"
    assert result.defined_names["phones"].attr_text == "'Ansprechpersonen'!$C$4:$C$8"
    root = WorkbookPackage(destination).read_sheet("Ansprechpersonen")
    validation_node = root.find("m:dataValidations/m:dataValidation", NS)
    assert validation_node is not None
    assert validation_node.get("sqref") == "C4:C8"
    assert validation_node.findtext(tag("formula1")) == "=$C$4:$C$5"
    conditional = root.find(tag("conditionalFormatting"))
    assert conditional is not None
    assert conditional.get("sqref") == "D4:D8"
    assert conditional.findtext("m:cfRule/m:formula", namespaces=NS) == 'C4<>""'
    _, table = WorkbookPackage(destination).table("contacts")
    assert [column.get("name") for column in table.findall("m:tableColumns/m:tableColumn", NS)] == [
        "id",
        "name",
        "phone",
        "display",
    ]
    assert group_input_columns(WorkbookPackage(destination)) == {}


def test_discontinuous_range_is_not_silently_expanded() -> None:
    with pytest.raises(ValueError, match="auseinanderfallen"):
        remap_reference("A1:B9", {1: 1, 2: 3, 3: 2})


def test_formula_remapping_keeps_literals_and_structured_fields() -> None:
    assert (
        remap_formula('IF(A4="D4",contacts[phone],$D$4)', "Test", {"Test": {4: 3}})
        == 'IF(A4="D4",contacts[phone],$C$4)'
    )


@pytest.mark.parametrize(
    ("expression", "expected"),
    [
        ('COUNTIF(INDEX($D:$D,ROW($D$3)+1):D4,"?*")', 'COUNTIF(INDEX($C:$C,ROW($C$3)+1):C4,"?*")'),
        ("D4:INDEX($D:$D,9)", "C4:INDEX($C:$C,9)"),
    ],
)
def test_dynamic_range_endpoints_move_with_their_column(expression: str, expected: str) -> None:
    assert remap_formula(expression, "Test", {"Test": {4: 3}}) == expected

"""Direkte Kommentare bleiben bei erledigten und erneut offenen Prüfungen erhalten."""

from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from data_imports.workbook_check_comments import enable_direct_check_comments
from data_imports.workbook_package import NS, WorkbookPackage, cell_map, serialize, set_cell, tag


def test_direct_comment_stays_with_check_when_counts_change(tmp_path: Path) -> None:
    path = tmp_path / "checks.xlsx"
    workbook = Workbook()
    view = workbook.active
    assert view is not None
    view.title = "Prüfungen"
    workbook.create_sheet("Anleitung")
    catalog = workbook.create_sheet("Prüfkatalog")
    catalog.append(["Katalog"])
    catalog.append(["Kopf"])
    catalog.append(
        [
            "Prüfung-ID",
            "Art",
            "Blatt",
            "Prüfung",
            "Erläuterung",
            "Anzahl",
            "Kundenkommentar",
            "Bearbeitung",
            "Status",
            "Rang",
        ]
    )
    catalog.append(["P001", "Fehler", "Kinder", "Erste", "Text", 0, "Kommentar zur erledigten Prüfung", "", "ok", 0])
    catalog.append(["P002", "Fehler", "Kinder", "Zweite", "Text", 1, "Kommentar zur offenen Prüfung", "", "offen", 1])
    catalog.add_table(Table(displayName="check_catalog", ref="A3:J5"))
    workbook.save(path)
    package = WorkbookPackage(path)
    styles = ET.fromstring(package.parts["xl/styles.xml"])
    formats = styles.find(tag("cellXfs"))
    assert formats is not None
    while len(formats) <= 100:
        formats.append(deepcopy(formats[0]))
    formats.set("count", str(len(formats)))
    package.parts["xl/styles.xml"] = serialize(styles)
    enable_direct_check_comments(package)
    cells = cell_map(package.read_sheet("Prüfungen"))
    assert package.value(cells["G7"]) == "Kommentar zur erledigten Prüfung"
    assert package.value(cells["G8"]) == "Kommentar zur offenen Prüfung"
    assert cells["G7"].find(tag("f")) is None
    assert "MATCH($H7,check_catalog[Prüfung-ID],0)" in cells["E7"].findtext(tag("f"), "")
    root = package.read_sheet("Prüfungen")
    set_cell(root, "G7", "Neuer Freitext")
    package.save_sheet("Prüfungen", root)
    root = package.read_sheet("Prüfkatalog")
    set_cell(root, "F4", "1")
    set_cell(root, "F5", "0")
    set_cell(root, "J4", "1")
    set_cell(root, "J5", "0")
    package.save_sheet("Prüfkatalog", root)
    enable_direct_check_comments(package)
    root = package.read_sheet("Prüfungen")
    cells = cell_map(root)
    assert package.value(cells["G7"]) == "Neuer Freitext"
    assert package.value(cells["H7"]) == "P001"
    assert package.value(cells["G8"]) == "Kommentar zur offenen Prüfung"
    assert root.find("m:conditionalFormatting/m:cfRule/m:formula", NS).text == "$E7=0"

"""Native XLSX-Patches erhalten Kundenwerte und unbekannte Excel-Metadaten."""

from copy import deepcopy
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table

from data_imports.workbook_package import (
    NS,
    WorkbookPackage,
    cell_map,
    expand_shared_formulas,
    merge_calculated_values,
    serialize,
    set_cell,
    tag,
)
from data_imports.workbook_upgrade import UpgradeReport, _repair


def test_native_edit_preserves_unknown_extensions_and_input(tmp_path: Path) -> None:
    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Kinder"
    sheet.append(["person_id", "notes"])
    sheet.append(["C9001", "=Literal"])
    sheet["B2"].data_type = "s"
    workbook.save(source)
    package = WorkbookPackage(source)
    root = package.read_sheet("Kinder")
    extension = ET.SubElement(ET.SubElement(root, tag("extLst")), tag("ext"), {"uri": "synthetic-extension"})
    ET.SubElement(extension, "{urn:synthetic:future}property", {"value": "keep"})
    before = deepcopy(extension)
    set_cell(root, "C2", 'IF(A2="","","berechnet")', formula=True)
    package.save_sheet("Kinder", root)
    destination = tmp_path / "edited.xlsx"
    package.write(destination)
    result = WorkbookPackage(destination)
    cells = cell_map(result.read_sheet("Kinder"))
    assert result.value(cells["B2"]) == "=Literal"
    assert ET.tostring(result.read_sheet("Kinder").find("m:extLst/m:ext", NS)) == ET.tostring(before)
    assert result.parts["xl/styles.xml"] == WorkbookPackage(source).parts["xl/styles.xml"]


def test_counters_no_longer_depend_on_deleted_predecessor(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Kinder"
    sheet.append(["Test"])
    sheet.append(["Eingaben"])
    sheet.append(["person_id", "issue_err", "issue_hint", "issue_kind", "issue_text", "issue_run"])
    for row in range(4, 7):
        sheet.cell(row, 1, f"C90{row}")
        sheet.cell(row, 2, f'=IF(A{row}="","","Befund")')
        sheet.cell(row, 3, f'=IF(A{row}="","","")')
        sheet.cell(row, 4, f'=IF(B{row}<>"","Fehler","")')
        sheet.cell(row, 5, f"=B{row}")
        sheet.cell(row, 6, f'=N(F{row - 1})+(D{row}<>"")')
    sheet["F5"] = '=N(#REF!)+(D5<>"")'
    sheet.add_table(Table(displayName="person", ref="A3:F6"))
    source = tmp_path / "counters.xlsx"
    workbook.save(source)
    package = WorkbookPackage(source)
    report = UpgradeReport(source_sha256="synthetic")
    _repair(package, "Kinder", ("person", "B", "C", "D", "E", "F"), report)
    cells = cell_map(package.read_sheet("Kinder"))
    assert package.value(cells["A5"]) == "C905"
    for row in range(4, 7):
        formula = cells[f"F{row}"].findtext(tag("f"), "")
        assert "#REF!" not in formula
        assert f"F{row - 1}" not in formula
        assert "ROW($D$3)" in formula
    _, table = package.table("person")
    assert all(
        column.find(tag("calculatedColumnFormula")) is not None
        for column in table.findall("m:tableColumns/m:tableColumn", NS)[1:]
    )


def test_cache_merge_only_updates_formula_values_and_rejects_errors(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "Original"
    sheet["B1"] = '=A1&" Kopie"'
    source = tmp_path / "native.xlsx"
    workbook.save(source)
    calculated = WorkbookPackage(source)
    root = calculated.read_sheet("Sheet")
    set_cell(root, "A1", "Must not overwrite original")
    cell = cell_map(root)["B1"]
    cell.set("t", "str")
    cached_value = cell.find(tag("v"))
    if cached_value is None:
        cached_value = ET.SubElement(cell, tag("v"))
    cached_value.text = "Original Kopie"
    calculated.save_sheet("Sheet", root)
    cache = tmp_path / "cache.xlsx"
    calculated.write(cache)
    assert merge_calculated_values(source, cache) == 1
    actual = load_workbook(source, data_only=True)
    assert actual["Sheet"]["A1"].value == "Original"
    assert actual["Sheet"]["B1"].value == "Original Kopie"
    actual.close()
    root = calculated.read_sheet("Sheet")
    cell_map(root)["B1"].set("t", "e")
    cell_map(root)["B1"].find(tag("v")).text = "#REF!"
    calculated.save_sheet("Sheet", root)
    calculated.write(cache)
    before = source.read_bytes()
    with pytest.raises(ValueError, match="Formelfehler"):
        merge_calculated_values(source, cache)
    assert source.read_bytes() == before


def test_ignorable_namespace_survives_serialization() -> None:
    original = b'<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:xr="urn:synthetic:unused" mc:Ignorable="xr"><sheetData/></worksheet>'
    result = serialize(ET.fromstring(original), b'<?xml version="1.0"?>' + original)
    assert b'xmlns:xr="urn:synthetic:unused"' in result
    ET.fromstring(result)


def test_shared_formula_survives_replacement_of_its_anchor() -> None:
    root = ET.fromstring(
        f'<worksheet xmlns="{NS["m"]}"><sheetData><row r="2"><c r="M2"><f t="shared" si="0" ref="M2:M10">M1+L1</f><v>2</v></c></row><row r="10"><c r="M10"><f t="shared" si="0"/><v>10</v></c></row></sheetData></worksheet>'
    )
    expand_shared_formulas(root)
    set_cell(root, "M2", "1", formula=True)
    follower = cell_map(root)["M10"].find(tag("f"))
    assert follower.text == "M9+L9"
    assert follower.get("t") is None


def test_review_rejects_orphan_shared_formula(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from data_imports.workbook_review import review_workbook
    from shared_modules.config import Config
    from tests.fixtures.create_masterdata import create_workbook

    monkeypatch.setattr(Config, "_instance", None)
    config = Config(Path(__file__).resolve().parents[1] / ".config/wegpiraten_config.yaml")
    source = create_workbook(tmp_path / "fixture.xlsx")
    package = WorkbookPackage(source)
    root = package.read_sheet("Berichte")
    cell = set_cell(root, "J4", "", formula=True)
    cell.find(tag("f")).attrib.update(t="shared", si="99")
    ET.SubElement(cell, tag("v")).text = "0"
    package.save_sheet("Berichte", root)
    package.write(source)
    result = review_workbook(config, source)
    assert any(issue.technical and issue.field == "J4" and "Sammelformel" in issue.message for issue in result.issues)


@pytest.mark.parametrize("valid_order", [False, True])
def test_review_checks_cell_format_child_order(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, valid_order: bool
) -> None:
    from data_imports.workbook_review import review_workbook
    from shared_modules.config import Config
    from tests.fixtures.create_masterdata import create_workbook

    monkeypatch.setattr(Config, "_instance", None)
    config = Config(Path(__file__).resolve().parents[1] / ".config/wegpiraten_config.yaml")
    source = create_workbook(tmp_path / "fixture.xlsx")
    package = WorkbookPackage(source)
    styles = ET.fromstring(package.parts["xl/styles.xml"])
    formats = styles.find(tag("cellXfs"))
    assert formats is not None
    style = ET.SubElement(formats, tag("xf"), {"numFmtId": "0", "fontId": "0", "fillId": "0", "borderId": "0"})
    children = ("alignment", "protection") if valid_order else ("protection", "alignment")
    for name in children:
        ET.SubElement(style, tag(name))
    formats.set("count", str(len(formats)))
    package.parts["xl/styles.xml"] = serialize(styles, package.parts["xl/styles.xml"])
    package.write(source)
    result = review_workbook(config, source)
    issues = [issue for issue in result.issues if issue.table == "styles.xml"]
    assert bool(issues) is not valid_order
    assert all(issue.technical for issue in issues)


def test_remove_empty_history_preserves_remaining_order_colors_and_local_names(tmp_path: Path) -> None:
    from openpyxl.workbook.defined_name import DefinedName

    source = tmp_path / "source.xlsx"
    workbook = Workbook()
    workbook.active.title = "Before"
    workbook.create_sheet("History")
    after = workbook.create_sheet("After")
    after.sheet_properties.tabColor = "008800"
    workbook.defined_names.add(DefinedName("local_check", attr_text="After!$A$1", localSheetId=2))
    workbook.save(source)
    package = WorkbookPackage(source)
    package.remove_sheet("History")
    destination = tmp_path / "result.xlsx"
    package.write(destination)
    result = WorkbookPackage(destination)
    assert list(result.sheets) == ["Before", "After"]
    assert result.workbook.find("m:definedNames/m:definedName", NS).get("localSheetId") == "1"
    assert result.read_sheet("After").find("m:sheetPr/m:tabColor", NS).get("rgb") == "00008800"


def test_review_detects_nonimported_rows_and_duplicate_keys(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from data_imports.workbook_review import review_workbook
    from shared_modules.config import Config
    from tests.fixtures.create_masterdata import create_workbook

    monkeypatch.setattr(Config, "_instance", None)
    config = Config(Path(__file__).resolve().parents[1] / ".config/wegpiraten_config.yaml")
    source = create_workbook(tmp_path / "fixture.xlsx")
    workbook = load_workbook(source)
    sheet = workbook["Berichte"]
    # Die zweite Zeile liegt ausserhalb der benannten Tabelle.
    sheet["A4"] = "R9001"
    sheet["B4"] = "A26092"
    sheet["C4"] = "A26092"
    workbook.save(source)
    workbook.close()
    result = review_workbook(config, source)
    assert any(issue.technical and "ausserhalb" in issue.message for issue in result.issues)
    assert not any(issue.table == "masterdata_contact_person" and issue.field == "gender" for issue in result.issues)
    workbook = load_workbook(source)
    workbook["Berichte"].tables["report"].ref = "A2:H4"
    workbook.save(source)
    workbook.close()
    result = review_workbook(config, source)
    assert any(issue.key == "R9001" and "doppelt" in issue.message for issue in result.issues)

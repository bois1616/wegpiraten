"""Prüft die synthetische Importfixture unabhängig vom Migrations-Build."""

from pathlib import Path

from openpyxl import load_workbook
from pytest import MonkeyPatch

from shared_modules.accordix import validate_coded_value
from tests.fixtures.create_masterdata import TABLES, create_workbook, fixture_records


def test_fixture_contains_import_tables_and_calculated_foreign_keys(tmp_path: Path) -> None:
    """Alle benötigten Tabellen und die beiden Auswahl-FKs sind direkt lesbar."""
    path = create_workbook(tmp_path / "masterdata.xlsx")
    workbook = load_workbook(path, data_only=True)
    actual = {name: table for sheet in workbook for name, table in sheet.tables.items()}
    assert set(actual) == set(TABLES)
    records = fixture_records()
    for name, (title, columns) in TABLES.items():
        sheet = workbook[title]
        assert tuple(c.value for c in sheet[2]) == columns
        assert sheet.max_row - 2 == len(records[name])
        assert all(c.data_type != "f" for row in sheet for c in row)
    assert "predecessor_mandate_id" in TABLES["mandate"][1]
    assert "mandate_id" in TABLES["report"][1]
    assert "masterdata_client" not in actual
    assert all(r["mandate_id"] != "SA" for r in records["mandate"])


def test_fixture_covers_family_parallel_renewal_exit_and_substitution() -> None:
    """Die Fälle verlangen unterschiedliche Index-, Zeit- und Zuordnungslogik."""
    records = fixture_records()
    family = [r for r in records["person"] if r.get("family_id") == "Testfamilie"]
    assert len(family) == 2
    assert len([r for r in records["mandate_person"] if r["person_id"] == "C9001"]) == 3
    assert any(r.get("predecessor_mandate_id") for r in records["mandate"])
    assert any(r.get("end_date") for r in records["mandate_person"])
    assert any(r["role"] == "S" for r in records["relation_mandate_emp"])
    assert any(not r["TS"] for r in records["masterdata_employee"])


def test_fixture_uses_valid_accordix_codes() -> None:
    """Die Ausgangsfixture enthält keine unbeabsichtigten Codefehler."""
    records = fixture_records()
    for row in records["person"]:
        for field in ("gender", "uma_umf", "canton_of_residence"):
            assert validate_coded_value(field, row[field]) is None
    for row in records["mandate_person"]:
        assert validate_coded_value("leaving_reason", row.get("leaving_reason")) is None


def test_new_config_models_map_fixture(tmp_path: Path, monkeypatch: MonkeyPatch) -> None:
    """Die vorbereiteten Modelle lesen alle Felder und wandeln Stunden einmal in Minuten."""
    from data_imports.import_masterdata import DEFAULT_TABLE_MAPPINGS, map_row, read_excel_table
    from shared_modules.config import Config

    monkeypatch.setattr(Config, "_instance", None)
    config = Config(Path(__file__).resolve().parents[1] / ".config" / "wegpiraten_config.yaml")
    path = create_workbook(tmp_path / "masterdata.xlsx")
    models = {
        "person": "person",
        "mandate": "mandate",
        "mandate_person": "mandate_person",
        "masterdata_contact_person": "contact_person",
        "relation_mandate_emp": "mandate_employee_relation",
        "report": "report",
    }
    for table, entity in models.items():
        fields = config.models[entity].fields
        frame = read_excel_table(path, table)
        assert all(field.excel_column in frame.columns for field in fields)
        mapped = map_row(
            frame.iloc[0], {field.excel_column: field for field in fields}, [field.name for field in fields]
        )
        if entity == "mandate":
            assert mapped["allowed_travel_time"] == 30
            assert mapped["allowed_direct_effort"] == 600
    assert "client" not in config.models
    assert "masterdata_client" not in DEFAULT_TABLE_MAPPINGS
    assert DEFAULT_TABLE_MAPPINGS["mandate"]["target"] == "mandate"
